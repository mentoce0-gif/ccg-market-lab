"""#0 の掲載用プレビュー画像を、納品するサンプル版の実物から作る（v2 §7.2「販売の説明」：プレビューと実物の一致）。

使い方：python make_previews.py   （LibreOffice Calc・pymupdf・Pillow が必要）
- 納品ファイルは変えない。コピーに印刷範囲を付け、LibreOffice で PDF にして画像にする。
- 画像に載せる文言は、ブックの中に実際にある機能だけにする。
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / "Teacher_Command_Center_Sample.xlsx"
OUT = HERE / "previews"
W, H = 2000, 1500   # 4:3 [未検証：Etsy の推奨比率]
NAVY, INK, PAPER, ACCENT = (27, 73, 94), (30, 30, 30), (247, 249, 250), (197, 80, 40)
FONT = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"

AREAS = {"Dashboard": "A1:L40", "Pacing": "A1:AA28", "Gradebook": "A1:V33", "Standings": "A1:R30",
         "Student Report": "A1:H40", "Start Here": "A1:B39", "Setup": "A1:H29"}

SLIDES = [
    ("01_dashboard.png", "Dashboard",
     "Teacher Command Center", "9 linked tabs: plan the term, grade it, see who needs help"),
    ("02_pacing.png", "Pacing",
     "Pacing plan sets the due dates", "Units by week on a timeline. Assignments are due at the end of their unit"),
    ("03_gradebook.png", "Gradebook",
     "Gradebook that knows what is due", "Type M for missing, EX for excused. Past-due blanks count as missing"),
    ("04_standings.png", "Standings",
     "Grade, turn-in rate and class standing", "Performance index with rank and percentile, plus an optional curved grade"),
    ("05_student_report.png", "Student Report",
     "One-page student report", "Pick a student: grade, missing work and every assignment's status"),
    ("06_setup.png", "Setup",
     "Your weights, your grade scale", "Category weights, letter-grade cut-offs and at-risk lines are editable"),
]


def render(tmp: Path, sheet: str) -> Image.Image:
    wb = load_workbook(SRC)
    for ws in wb.worksheets:
        ws.sheet_state = "visible" if ws.title == sheet else "hidden"
    wb.active = wb.sheetnames.index(sheet)
    ws = wb[sheet]
    ws.print_area = AREAS[sheet]
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 0.2
    x = tmp / f"{sheet.replace(' ', '_')}.xlsx"
    wb.save(x)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/lo", "--headless", "--convert-to", "pdf",
                    "--outdir", str(tmp), str(x)], check=True, capture_output=True, timeout=180)
    pix = pymupdf.open(x.with_suffix(".pdf"))[0].get_pixmap(dpi=220)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    bbox = ImageChops.difference(img, Image.new("RGB", img.size, "white")).getbbox()
    return img.crop(bbox) if bbox else img


def slide(shot: Image.Image, title: str, sub: str) -> Image.Image:
    im = Image.new("RGB", (W, H), PAPER)
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, W, 230], fill=NAVY)
    d.text((80, 50), title, font=ImageFont.truetype(BOLD, 78), fill="white")
    d.text((80, 150), sub, font=ImageFont.truetype(FONT, 44), fill=(214, 228, 235))
    box_w, box_h = W - 120, H - 230 - 110
    s = shot.copy()
    s.thumbnail((box_w, box_h), Image.LANCZOS)
    x, y = (W - s.width) // 2, 230 + 50 + (box_h - s.height) // 2
    d.rectangle([x - 6, y - 6, x + s.width + 6, y + s.height + 6], fill=(210, 216, 220))
    im.paste(s, (x, y))
    d.text((80, H - 60), "Screenshot of the included sample file (fictional class). Rendered in LibreOffice.",
           font=ImageFont.truetype(FONT, 30), fill=(110, 110, 110))
    return im


def included() -> Image.Image:
    im = Image.new("RGB", (W, H), PAPER)
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, W, 230], fill=NAVY)
    d.text((80, 50), "What you get", font=ImageFont.truetype(BOLD, 78), fill="white")
    d.text((80, 150), "Instant digital download - 2 spreadsheet files", font=ImageFont.truetype(FONT, 44), fill=(214, 228, 235))
    lines = [
        ("Blank file", "ready for your class"),
        ("Sample file", "a fictional class to explore first"),
        ("9 tabs", "Start Here, Dashboard, Setup, Roster, Pacing,"),
        ("", "Assignments, Gradebook, Standings, Student Report"),
        ("Capacity", "40 students, 60 assignments, 20 units per file"),
        ("Format", ".xlsx for Excel; opens in Google Sheets"),
        ("Set up", "5 steps on the Start Here tab"),
    ]
    f1, f2 = ImageFont.truetype(BOLD, 46), ImageFont.truetype(FONT, 46)
    y = 300
    for k, v in lines:
        d.text((90, y), k, font=f1, fill=ACCENT)
        d.text((420, y), v, font=f2, fill=INK)
        y += 84
    d.text((90, H - 90), "Digital download only. No physical item is shipped. Made with AI assistance.",
           font=ImageFont.truetype(FONT, 32), fill=(90, 90, 90))
    return im


def main():
    if not shutil.which("soffice"):
        sys.exit("LibreOffice (soffice) が必要")
    OUT.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        for fname, sheet, title, sub in SLIDES:
            slide(render(tmp, sheet), title, sub).save(OUT / fname, optimize=True)
            print(OUT / fname)
        included().save(OUT / "07_whats_included.png", optimize=True)
        print(OUT / "07_whats_included.png")


if __name__ == "__main__":
    main()
