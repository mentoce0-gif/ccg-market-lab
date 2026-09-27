"""Teacher Command Center: pacing plan -> assignments -> roster x gradebook -> metrics -> standings -> dashboard.

Excel / Google Sheets / LibreOffice compatible (only Excel-2007-era functions).
Usage: python build_command_center.py OUTDIR
"""
import random
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

N_STU, N_ASG, N_UNIT, N_WEEK = 40, 60, 20, 20
# Gradebook layout
GB_FIRST = 10
GB_LAST = GB_FIRST + N_STU - 1          # 49
SC1 = 13                                 # M
SC2 = SC1 + N_ASG - 1                    # BT
S1, S2 = L(SC1), L(SC2)
# Assignments layout
A_FIRST, A_LAST = 6, 6 + N_ASG - 1       # 6..65
# Pacing layout
P_FIRST, P_LAST = 6, 6 + N_UNIT - 1      # 6..25
WK1 = 8                                  # H
WKN = WK1 + N_WEEK - 1                   # AA
# Roster
R_FIRST, R_LAST = 6, 6 + N_STU - 1
# Standings
ST_FIRST, ST_LAST = 8, 8 + N_STU - 1     # 8..47

CATS = [("Homework", 0.20), ("Quizzes", 0.20), ("Tests", 0.40), ("Projects", 0.10), ("Participation", 0.10)]
SCALE = [(0.00, "F"), (0.60, "D-"), (0.63, "D"), (0.67, "D+"), (0.70, "C-"), (0.73, "C"), (0.77, "C+"),
         (0.80, "B-"), (0.83, "B"), (0.87, "B+"), (0.90, "A-"), (0.93, "A")]
CURVE = [(0.00, "F"), (0.05, "D"), (0.20, "C"), (0.55, "B"), (0.85, "A")]

FONT = "Arial"
INK, MUTED = "1E2B33", "5E6C75"
TEAL, TEAL_LIGHT, TEAL_PALE = "17495A", "D6E6EA", "EEF5F6"
ACCENT = "C8501A"
CALC = PatternFill("solid", fgColor="F1F3F4")
WHITE = PatternFill("solid", fgColor="FFFFFF")
HEAD = PatternFill("solid", fgColor=TEAL)
SUB = PatternFill("solid", fgColor=TEAL_LIGHT)
SIDE = Side(style="thin", color="C5D0D4")
BOX = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)


def ft(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def banner(ws, title, subtitle, width_cols):
    ws["A1"] = title
    ws["A1"].font = ft(18, True, "FFFFFF")
    ws["A2"] = subtitle
    ws["A2"].font = ft(9, color="DCEBEF")
    for c in range(1, width_cols + 1):
        ws.cell(row=1, column=c).fill = HEAD
        ws.cell(row=2, column=c).fill = HEAD
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18


def hdr(cell, text, size=10):
    cell.value = text
    cell.font = ft(size, True, "FFFFFF")
    cell.fill = HEAD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BOX


def inp(cell, value=None, fmt=None, bold=False):
    if value is not None:
        cell.value = value
    cell.fill = WHITE
    cell.border = BOX
    cell.font = ft(10, bold)
    if fmt:
        cell.number_format = fmt


def calc(cell, value, fmt=None, bold=False, center=False, size=10):
    cell.value = value
    cell.fill = CALC
    cell.border = BOX
    cell.font = ft(size, bold)
    if fmt:
        cell.number_format = fmt
    if center:
        cell.alignment = Alignment(horizontal="center")


def note(cell, text):
    cell.value = text
    cell.font = ft(9, italic=True, color=MUTED)


# ---------------------------------------------------------------- sample data
def sample(seed=20261005, n_students=24):
    rnd = random.Random(seed)
    term_start = date(2026, 8, 17)
    units = [("Foundations", 1, 3), ("Linear Equations", 4, 3), ("Functions", 7, 3), ("Midterm Review", 10, 1),
             ("Systems of Equations", 11, 3), ("Quadratics", 14, 3), ("Final Review", 17, 2)]
    asg = []
    for u, (uname, sw, wk) in enumerate(units):
        if uname.endswith("Review"):
            asg.append((f"{uname} Packet", uname, "Homework", 20, None))
            asg.append((f"{uname.split()[0]} Exam", uname, "Tests", 100, None))
            continue
        base = term_start + timedelta(days=(sw - 1) * 7)
        asg.append((f"{uname}: HW A", uname, "Homework", 10, base + timedelta(days=3)))
        asg.append((f"{uname}: HW B", uname, "Homework", 10, base + timedelta(days=10)))
        asg.append((f"{uname}: Quiz", uname, "Quizzes", 25, base + timedelta(days=11)))
        if uname == "Functions":
            asg.append(("Functions Project", uname, "Projects", 50, None))
        asg.append((f"{uname}: Participation", uname, "Participation", 10, None))
        asg.append((f"{uname}: Test", uname, "Tests", 100, None))
    first = ["Ava", "Liam", "Mia", "Noah", "Emma", "Lucas", "Sofia", "Ethan", "Isabella", "Mason", "Chloe", "Logan",
             "Harper", "Elijah", "Amelia", "James", "Evelyn", "Benjamin", "Abigail", "Henry", "Ella", "Jack", "Aria", "Owen"]
    last = ["Johnson", "Smith", "Garcia", "Brown", "Davis", "Miller", "Wilson", "Moore", "Taylor", "Anderson",
            "Thomas", "Martinez", "Lee", "Clark", "Lewis", "Walker", "Hall", "Young", "King", "Wright", "Scott",
            "Green", "Adams", "Baker"]
    names = [f"{first[i]} {last[i]}" for i in range(n_students)]
    profile = [(rnd.uniform(0.72, 0.99), rnd.uniform(0.88, 1.0)) for _ in names]
    return term_start, units, asg, names, profile, rnd


def unit_end(term_start, sw, wk):
    return term_start + timedelta(days=(sw + wk - 2) * 7 + 4)


def fill_scores(term_start, units, asg, names, profile, rnd, asof):
    uend = {u: unit_end(term_start, sw, wk) for u, sw, wk in units}
    scores = {}
    for si, (skill, diligence) in enumerate(profile):
        for ai, (_, unit, cat, pts, due) in enumerate(asg):
            d = due or uend[unit]
            if d > asof:
                continue
            r = rnd.random()
            if r > diligence:
                scores[(si, ai)] = "M" if rnd.random() < 0.7 else None   # missing: marked or left blank
                continue
            if rnd.random() < 0.03:
                scores[(si, ai)] = "EX"
                continue
            val = max(0.0, min(1.0, rnd.gauss(skill, 0.07)))
            scores[(si, ai)] = round(val * pts) if pts >= 20 else round(val * pts, 1)
    return scores


# ---------------------------------------------------------------- build
def build(path, term_start, units, asg, names, scores, asof_value="=TODAY()", class_name="Period 3 - Algebra I",
          teacher="Ms. Rivera"):
    wb = Workbook()

    # ====================== Start Here
    ws = wb.active
    ws.title = "Start Here"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 104
    banner(ws, "Teacher Command Center", "Plan the term, grade it, and see who needs help - all tabs work together.", 2)
    lines = [
        ("", None),
        ("What is linked to what", ft(13, True)),
        ("Pacing plan  >  Assignments get their due dates  >  Gradebook knows what is due  >  submission rate, missing work,", ft(11)),
        ("grades and class standing update by themselves  >  Dashboard and Student Report show the result.", ft(11)),
        ("", None),
        ("Set up in 15 minutes", ft(13, True)),
        ("1.  Setup: class name, the Monday your term starts, category weights and grade scale.", ft(11)),
        ("2.  Roster: type your students (up to 40).", ft(11)),
        ("3.  Pacing: list your units with a start week and length. The timeline colors itself.", ft(11)),
        ("4.  Assignments: pick the unit for each assignment. It is due at the end of that unit unless you type a date.", ft(11)),
        ("5.  Gradebook: enter scores. Type M for missing, EX for excused. Past-due blanks count as missing.", ft(11)),
        ("", None),
        ("What you get", ft(13, True)),
        ("Dashboard: class average, submission rate, students at risk, grade spread, workload per week.", ft(11)),
        ("Standings: each student's grade, submission rate and a combined performance index with rank and percentile,", ft(11)),
        ("plus an optional curved grade if your school allows relative grading.", ft(11)),
        ("Student Report: a one-page summary for conferences, with every assignment marked Submitted, Missing, Excused or Not due yet.", ft(11)),
        ("", None),
        ("Cell colors", ft(13, True)),
        ("White cells: type here.   Gray cells: calculated for you. Please do not type over them.", ft(11)),
        ("", None),
        ("How the numbers work", ft(13, True)),
        ("Category % = points earned / points possible on graded work. M counts as 0. EX is skipped. Blank work that is past due", ft(11)),
        ("counts as 0 when Setup says Yes. The overall grade weights the categories that have graded work so far.", ft(11)),
        ("Submission rate = work turned in / work that is due (EX not counted). Performance index blends the overall grade and the", ft(11)),
        ("submission rate with the weights on Setup. Percentile compares each student with the rest of the class.", ft(11)),
        ("Extra credit: give an assignment 0 points possible.", ft(11)),
        ("", None),
        ("Google Sheets", ft(13, True)),
        ("Upload the file to Google Drive, open it, then choose File > Save as Google Sheets.", ft(11)),
        ("", None),
        ("FAQ", ft(13, True)),
        ("A student shows Missing work I have not graded yet?  Set the date on Setup (As of) or type the score. Blank + past due = missing.", ft(11)),
        ("I do not use one of the categories?  Set its weight to 0%. The rest still adds up.", ft(11)),
        ("More than 40 students or 60 assignments?  Make a copy of the file for each class period.", ft(11)),
        ("", None),
        ("Version 2.0", ft(9, color=MUTED)),
    ]
    for i, (text, font) in enumerate(lines, start=3):
        c = ws.cell(row=i, column=2, value=text or None)
        if font:
            c.font = font

    # ====================== Setup
    st = wb.create_sheet("Setup")
    st.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [30, 16, 3, 10, 10, 3, 12, 12]):
        st.column_dimensions[col].width = w
    banner(st, "Setup", "Fill the white cells once at the start of the term.", 8)
    rows = [(4, "Class", class_name, None), (5, "Teacher", teacher, None), (6, "Term starts (Monday)", term_start, "yyyy-mm-dd"),
            (7, "Weeks in the term", 18, "0"), (8, "As of (leave =TODAY() or type a date)", asof_value, "yyyy-mm-dd"),
            (9, "Count past-due blanks as 0?", "Yes", None)]
    for r, lab, val, fmt in rows:
        st.cell(row=r, column=1, value=lab).font = ft(10, True)
        inp(st.cell(row=r, column=2), val, fmt)
    dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    st.add_data_validation(dv_yes)
    dv_yes.add("B9")

    hdr(st["A11"], "Category")
    hdr(st["B11"], "Weight")
    for i, (c, w) in enumerate(CATS):
        inp(st.cell(row=12 + i, column=1), c)
        inp(st.cell(row=12 + i, column=2), w, "0%")
    st["A17"] = "Total"
    st["A17"].font = ft(10, True)
    calc(st["B17"], "=SUM(B12:B16)", "0%", True)
    st["A18"] = '=IF(ROUND(B17,4)=1,"Weights add up to 100%.","Check: weights add up to "&TEXT(B17,"0%")&".")'
    st["A18"].font = ft(9, italic=True, color=MUTED)

    hdr(st["D11"], "Min %")
    hdr(st["E11"], "Grade")
    for i, (mn, g) in enumerate(SCALE):
        inp(st.cell(row=12 + i, column=4), mn, "0%")
        inp(st.cell(row=12 + i, column=5), g)
    note(st["D24"], "Lowest to highest.")

    st["A21"] = "Performance index"
    st["A21"].font = ft(11, True)
    st["A22"] = "Weight of overall grade"
    st["A23"] = "Weight of submission rate"
    inp(st["B22"], 0.70, "0%")
    inp(st["B23"], 0.30, "0%")
    st["A26"] = "At-risk flags"
    st["A26"].font = ft(11, True)
    st["A27"] = "Overall grade below"
    st["A28"] = "Submission rate below"
    inp(st["B27"], 0.70, "0%")
    inp(st["B28"], 0.80, "0%")
    note(st["A29"], "'Watch' = within 10 points of either line.")

    hdr(st["G11"], "Percentile from")
    hdr(st["H11"], "Curved grade")
    for i, (p, g) in enumerate(CURVE):
        inp(st.cell(row=12 + i, column=7), p, "0%")
        inp(st.cell(row=12 + i, column=8), g)
    note(st["G17"], "Optional relative grading.")
    note(st["G18"], "Default: top 15% A, next 30% B,")
    note(st["G19"], "next 35% C, next 15% D, last 5% F.")

    # ====================== Roster
    ro = wb.create_sheet("Roster")
    ro.sheet_view.showGridLines = False
    for col, w in zip("ABCD", [5, 28, 14, 40]):
        ro.column_dimensions[col].width = w
    banner(ro, "Roster", "One row per student. Names flow into every other tab.", 4)
    for cell, t in zip(["A5", "B5", "C5", "D5"], ["#", "Student", "Student ID", "Notes"]):
        hdr(ro[cell], t)
    for i in range(N_STU):
        r = R_FIRST + i
        ro.cell(row=r, column=1, value=i + 1).font = ft(9, color=MUTED)
        for c in (2, 3, 4):
            inp(ro.cell(row=r, column=c))
    for i, n in enumerate(names):
        ro.cell(row=R_FIRST + i, column=2, value=n)
        ro.cell(row=R_FIRST + i, column=3, value=f"S{1001 + i}" if len(names) > 1 else "")
    ro.freeze_panes = "B6"

    # ====================== Pacing
    pc = wb.create_sheet("Pacing")
    pc.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [4, 24, 7, 7, 11, 11, 7]):
        pc.column_dimensions[col].width = w
    for c in range(WK1, WKN + 1):
        pc.column_dimensions[L(c)].width = 4.2
    banner(pc, "Pacing Plan", "Units by week. Assignments take their due date from the end of their unit.", WKN)
    pc["A3"] = "This week:"
    pc["A3"].font = ft(9, True, MUTED)
    pc["C3"] = '=IF(Setup!B8="","",INT((Setup!B8-Setup!B6)/7)+1)'
    pc["C3"].font = ft(10, True, ACCENT)
    for cell, t in zip(["A5", "B5", "C5", "D5", "E5", "F5", "G5"], ["#", "Unit", "Start wk", "Weeks", "Starts", "Ends", "Items"]):
        hdr(pc[cell], t, 9)
    for w in range(N_WEEK):
        c = WK1 + w
        n = pc.cell(row=4, column=c, value=w + 1)
        n.font = ft(8, True, MUTED)
        n.alignment = Alignment(horizontal="center")
        d = pc.cell(row=5, column=c, value=f"=Setup!$B$6+({L(c)}$4-1)*7")
        d.number_format = "m/d"
        d.font = ft(7, True, "FFFFFF")
        d.fill = HEAD
        d.alignment = Alignment(horizontal="center", text_rotation=90)
        d.border = BOX
    pc.row_dimensions[5].height = 34
    for i in range(N_UNIT):
        r = P_FIRST + i
        pc.cell(row=r, column=1, value=i + 1).font = ft(9, color=MUTED)
        inp(pc.cell(row=r, column=2))
        inp(pc.cell(row=r, column=3), None, "0")
        inp(pc.cell(row=r, column=4), None, "0")
        calc(pc.cell(row=r, column=5), f'=IF(OR(B{r}="",C{r}=""),"",Setup!$B$6+(C{r}-1)*7)', "m/d")
        calc(pc.cell(row=r, column=6), f'=IF(OR(B{r}="",C{r}=""),"",Setup!$B$6+(C{r}+MAX(1,D{r})-2)*7+4)', "m/d")
        calc(pc.cell(row=r, column=7), f'=IF(B{r}="","",COUNTIF(Assignments!$C${A_FIRST}:$C${A_LAST},B{r}))', "0", center=True)
        for c in range(WK1, WKN + 1):
            cell = pc.cell(row=r, column=c)
            cell.border = Border(left=Side(style="hair", color="D9E1E4"), right=Side(style="hair", color="D9E1E4"),
                                 top=SIDE, bottom=SIDE)
    for i, (u, sw, wk) in enumerate(units):
        r = P_FIRST + i
        pc.cell(row=r, column=2, value=u)
        pc.cell(row=r, column=3, value=sw)
        pc.cell(row=r, column=4, value=wk)
    grid = f"{L(WK1)}{P_FIRST}:{L(WKN)}{P_LAST}"
    pc.conditional_formatting.add(grid, FormulaRule(
        formula=[f"AND(ISNUMBER($C{P_FIRST}),{L(WK1)}$4>=$C{P_FIRST},{L(WK1)}$4<=$C{P_FIRST}+MAX(1,$D{P_FIRST})-1,{L(WK1)}$4=$C$3)"],
        fill=PatternFill("solid", fgColor=ACCENT), stopIfTrue=True))
    pc.conditional_formatting.add(grid, FormulaRule(
        formula=[f"AND(ISNUMBER($C{P_FIRST}),{L(WK1)}$4>=$C{P_FIRST},{L(WK1)}$4<=$C{P_FIRST}+MAX(1,$D{P_FIRST})-1)"],
        fill=PatternFill("solid", fgColor="3F7F92")))
    pc.conditional_formatting.add(f"{L(WK1)}4:{L(WKN)}4", FormulaRule(
        formula=[f"{L(WK1)}$4=$C$3"], fill=PatternFill("solid", fgColor=ACCENT), font=Font(color="FFFFFF", bold=True)))
    wr = P_LAST + 2
    pc.cell(row=wr, column=2, value="Assignments due").font = ft(10, True)
    pc.cell(row=wr + 1, column=2, value="Points due").font = ft(10, True)
    for c in range(WK1, WKN + 1):
        calc(pc.cell(row=wr, column=c), f"=COUNTIF(Assignments!$H${A_FIRST}:$H${A_LAST},{L(c)}$4)", '0;-0;""', center=True, size=9)
        calc(pc.cell(row=wr + 1, column=c), f"=SUMIF(Assignments!$H${A_FIRST}:$H${A_LAST},{L(c)}$4,Assignments!$E${A_FIRST}:$E${A_LAST})",
             '0;-0;""', center=True, size=8)
    pc.conditional_formatting.add(f"{L(WK1)}{wr}:{L(WKN)}{wr}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                                             color="9CC3CF"))
    pc.freeze_panes = f"{L(WK1)}6"

    # ====================== Assignments
    a = wb.create_sheet("Assignments")
    a.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKL", [4, 30, 20, 14, 8, 12, 11, 6, 6, 10, 10, 8]):
        a.column_dimensions[col].width = w
    banner(a, "Assignments", "Pick a unit and the due date fills in. Type a date only if it differs from the end of the unit.", 12)
    heads = ["#", "Assignment", "Unit", "Category", "Points", "Due date (optional)", "Due", "Week", "Due?",
             "Class avg", "Turned in", "Missing"]
    for i, t in enumerate(heads):
        hdr(a.cell(row=5, column=i + 1), t, 9)
    for i in range(N_ASG):
        r = A_FIRST + i
        gcol = L(SC1 + i)
        col_rng = f"Gradebook!{gcol}${GB_FIRST}:{gcol}${GB_LAST}"
        a.cell(row=r, column=1, value=i + 1).font = ft(9, color=MUTED)
        inp(a.cell(row=r, column=2))
        inp(a.cell(row=r, column=3))
        inp(a.cell(row=r, column=4))
        inp(a.cell(row=r, column=5), None, "0")
        inp(a.cell(row=r, column=6), None, "yyyy-mm-dd")
        calc(a.cell(row=r, column=7),
             f'=IF(B{r}="","",IF(F{r}<>"",F{r},IFERROR(INDEX(Pacing!$F${P_FIRST}:$F${P_LAST},MATCH(C{r},Pacing!$B${P_FIRST}:$B${P_LAST},0)),"")))',
             "m/d")
        calc(a.cell(row=r, column=8), f'=IF(G{r}="","",INT((G{r}-Setup!$B$6)/7)+1)', "0", center=True)
        calc(a.cell(row=r, column=9), f'=IF(AND(B{r}<>"",G{r}<>""),IF(G{r}<=Setup!$B$8,1,0),0)', '"yes";;"-"', center=True)
        calc(a.cell(row=r, column=10), f'=IF(OR(B{r}="",N(E{r})=0),"",IFERROR(AVERAGE({col_rng})/E{r},""))', "0%", center=True)
        calc(a.cell(row=r, column=11),
             f'=IF(I{r}<>1,"",IFERROR(COUNT({col_rng})/(SUMPRODUCT(--(Gradebook!$B${GB_FIRST}:$B${GB_LAST}<>""))-COUNTIF({col_rng},"EX")),""))',
             "0%", center=True)
        calc(a.cell(row=r, column=12),
             f'=IF(B{r}="","",IF(I{r}=1,SUMPRODUCT((Gradebook!$B${GB_FIRST}:$B${GB_LAST}<>"")*({col_rng}="")),0)+COUNTIF({col_rng},"M"))',
             '0;-0;""', center=True)
    for i, (name, unit, cat, pts, due) in enumerate(asg):
        r = A_FIRST + i
        a.cell(row=r, column=2, value=name)
        a.cell(row=r, column=3, value=unit)
        a.cell(row=r, column=4, value=cat)
        a.cell(row=r, column=5, value=pts)
        if due:
            a.cell(row=r, column=6, value=due)
    dv_u = DataValidation(type="list", formula1=f"=Pacing!$B${P_FIRST}:$B${P_LAST}", allow_blank=True)
    dv_c = DataValidation(type="list", formula1="=Setup!$A$12:$A$16", allow_blank=True)
    a.add_data_validation(dv_u)
    a.add_data_validation(dv_c)
    dv_u.add(f"C{A_FIRST}:C{A_LAST}")
    dv_c.add(f"D{A_FIRST}:D{A_LAST}")
    a["K3"] = "=Setup!B28"
    a["K3"].number_format = "0%"
    a["K3"].font = ft(8, color=MUTED)
    a["J3"] = "Flag below:"
    a["J3"].font = ft(8, True, MUTED)
    a.conditional_formatting.add(f"K{A_FIRST}:K{A_LAST}", FormulaRule(
        formula=[f'AND(ISNUMBER(K{A_FIRST}),K{A_FIRST}<$K$3)'], font=Font(color="9A2A1C", bold=True)))
    a.freeze_panes = "C6"

    # ====================== Gradebook
    g = wb.create_sheet("Gradebook")
    g.sheet_view.showGridLines = False
    widths = [4, 22, 8, 6, 8, 7, 8, 8, 8, 8, 8, 8]
    for i, w in enumerate(widths):
        g.column_dimensions[L(i + 1)].width = w
    for c in range(SC1, SC2 + 1):
        g.column_dimensions[L(c)].width = 6
    for c in range(1, SC2 + 1):
        g.cell(row=1, column=c).fill = HEAD
        g.cell(row=2, column=c).fill = HEAD
    g["A1"] = '="Gradebook  |  "&Setup!B4'
    g["A1"].font = ft(16, True, "FFFFFF")
    g["A2"] = "Enter scores in the white grid.  M = missing (0).  EX = excused.  Blank + past due = missing."
    g["A2"].font = ft(9, color="DCEBEF")
    g.row_dimensions[1].height = 28
    g["B3"] = "At-risk lines:"
    g["B3"].font = ft(8, True, MUTED)
    g["B3"].alignment = Alignment(horizontal="right")
    g["C3"] = "=Setup!B27"
    g["E3"] = "=Setup!B28"
    for cc in ("C3", "E3"):
        g[cc].number_format = "0%"
        g[cc].font = ft(8, color=MUTED)
        g[cc].alignment = Alignment(horizontal="center")
    g["G3"] = "Weight"
    g["G3"].font = ft(8, True, MUTED)
    g["G3"].alignment = Alignment(horizontal="right")
    for j in range(5):
        c = g.cell(row=3, column=8 + j, value=f"=Setup!B{12 + j}")
        c.number_format = "0%"
        c.font = ft(8, color=MUTED)
        c.alignment = Alignment(horizontal="center")
    for rr, lab in [(5, "Assignment >"), (6, "Category >"), (7, "Points >"), (8, "Due >")]:
        c = g.cell(row=rr, column=12, value=lab)
        c.font = ft(8, True, MUTED)
        c.alignment = Alignment(horizontal="right")
    for i in range(N_ASG):
        c = SC1 + i
        col = L(c)
        n = g.cell(row=4, column=c, value=i + 1)
        n.font = ft(7, color=MUTED)
        n.alignment = Alignment(horizontal="center")
        nm = g.cell(row=5, column=c, value=f'=IF(INDEX(Assignments!$B${A_FIRST}:$B${A_LAST},{col}$4)="","",INDEX(Assignments!$B${A_FIRST}:$B${A_LAST},{col}$4))')
        nm.font = ft(8, True)
        nm.fill = SUB
        nm.alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom", wrap_text=True)
        ct = g.cell(row=6, column=c, value=f'=IF({col}$5="","",INDEX(Assignments!$D${A_FIRST}:$D${A_LAST},{col}$4))')
        ct.font = ft(6, color=MUTED)
        ct.fill = SUB
        ct.alignment = Alignment(horizontal="center", shrink_to_fit=True)
        pt = g.cell(row=7, column=c, value=f'=IF({col}$5="",0,N(INDEX(Assignments!$E${A_FIRST}:$E${A_LAST},{col}$4)))')
        pt.number_format = '0;-0;""'
        pt.font = ft(8, True)
        pt.fill = SUB
        pt.alignment = Alignment(horizontal="center")
        du = g.cell(row=8, column=c, value=f"=INDEX(Assignments!$I${A_FIRST}:$I${A_LAST},{col}$4)")
        du.number_format = '"due";;""'
        du.font = ft(7, color=ACCENT)
        du.fill = SUB
        du.alignment = Alignment(horizontal="center")
        for rr in range(5, 9):
            g.cell(row=rr, column=c).border = BOX
    g.row_dimensions[5].height = 118
    heads = ["#", "Student", "Overall", "Grade", "Turned in", "Missing", "Score % (turned in)"]
    for i, t in enumerate(heads):
        hdr(g.cell(row=9, column=i + 1), t, 8)
    for j in range(5):
        c = g.cell(row=9, column=8 + j, value=f"=Setup!A{12 + j}")
        c.font = ft(8, True, "FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BOX
    for c in range(SC1, SC2 + 1):
        hdr(g.cell(row=9, column=c), "Score", 7)
    g.row_dimensions[9].height = 30

    for i in range(N_STU):
        r = GB_FIRST + i
        sc = f"${S1}{r}:${S2}{r}"
        g.cell(row=r, column=1, value=i + 1).font = ft(8, color=MUTED)
        calc(g.cell(row=r, column=2), f'=IF(Roster!B{R_FIRST + i}="","",Roster!B{R_FIRST + i})', bold=True)
        calc(g.cell(row=r, column=3),
             f'=IF($B{r}="","",IFERROR(SUMPRODUCT($H$3:$L$3,H{r}:L{r})/SUMPRODUCT($H$3:$L$3,--ISNUMBER(H{r}:L{r})),""))',
             "0.0%", bold=True)
        calc(g.cell(row=r, column=4), f'=IF(C{r}="","",VLOOKUP(C{r},Setup!$D$12:$E$23,2,TRUE))', bold=True, center=True)
        calc(g.cell(row=r, column=5),
             f'=IF($B{r}="","",IFERROR(SUMPRODUCT(${S1}$8:${S2}$8*ISNUMBER({sc}))/SUMPRODUCT(${S1}$8:${S2}$8*({sc}<>"EX")),""))',
             "0%", center=True)
        calc(g.cell(row=r, column=6), f'=IF($B{r}="","",SUMPRODUCT(${S1}$8:${S2}$8*({sc}=""))+COUNTIF({sc},"M"))', "0", center=True)
        calc(g.cell(row=r, column=7), f'=IF($B{r}="","",IFERROR(SUM({sc})/SUMPRODUCT(${S1}$7:${S2}$7*ISNUMBER({sc})),""))',
             "0%", center=True)
        for j in range(5):
            cc = L(8 + j)
            earned = f"SUMIFS({sc},${S1}$6:${S2}$6,{cc}$9)"
            possible = (f'SUMPRODUCT((${S1}$6:${S2}$6={cc}$9)*${S1}$7:${S2}$7*(ISNUMBER({sc})+({sc}="M")'
                        f'+(Setup!$B$9="Yes")*${S1}$8:${S2}$8*({sc}="")))')
            calc(g.cell(row=r, column=8 + j), f'=IF($B{r}="","",IFERROR({earned}/{possible},""))', "0%", center=True, size=9)
        for c in range(SC1, SC2 + 1):
            cell = g.cell(row=r, column=c)
            cell.fill = WHITE
            cell.border = BOX
            cell.font = ft(9)
            cell.alignment = Alignment(horizontal="center")
    for (si, ai), v in scores.items():
        if v is not None:
            g.cell(row=GB_FIRST + si, column=SC1 + ai, value=v)
    g.freeze_panes = f"{S1}{GB_FIRST}"
    gridr = f"{S1}{GB_FIRST}:{S2}{GB_LAST}"
    g.conditional_formatting.add(gridr, FormulaRule(formula=[f'UPPER({S1}{GB_FIRST})="M"'],
                                                    fill=PatternFill("solid", fgColor="FBE0C6"), font=Font(color="8A3B00", bold=True)))
    g.conditional_formatting.add(gridr, FormulaRule(formula=[f'UPPER({S1}{GB_FIRST})="EX"'], font=Font(color="8A949B", italic=True)))
    g.conditional_formatting.add(gridr, FormulaRule(
        formula=[f'AND({S1}{GB_FIRST}="",{S1}$8=1,$B{GB_FIRST}<>"")'], fill=PatternFill("solid", fgColor="FDF1E7")))
    g.conditional_formatting.add(f"C{GB_FIRST}:D{GB_LAST}", FormulaRule(
        formula=[f"AND(ISNUMBER($C{GB_FIRST}),$C{GB_FIRST}<$C$3)"],
        fill=PatternFill("solid", fgColor="F6D5D0"), font=Font(color="9A2A1C", bold=True)))
    g.conditional_formatting.add(f"E{GB_FIRST}:E{GB_LAST}", FormulaRule(
        formula=[f"AND(ISNUMBER($E{GB_FIRST}),$E{GB_FIRST}<$E$3)"],
        fill=PatternFill("solid", fgColor="F6D5D0"), font=Font(color="9A2A1C", bold=True)))

    # ====================== Standings
    sd = wb.create_sheet("Standings")
    sd.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKLMNOPQR", [4, 22, 9, 10, 10, 11, 7, 11, 9, 9, 11, 3, 3, 6, 22, 11, 9, 10]):
        sd.column_dimensions[col].width = w
    banner(sd, "Standings", "Relative view of the class: grade and turn-in rate combined, with rank and percentile.", 18)
    note(sd["A4"], "Performance index = overall grade and submission rate, weighted as on Setup. Rank 1 = highest.")
    note(sd["A5"], "Curved grade is optional: it assigns letters by percentile (see Setup). Use it only if your school allows relative grading.")
    heads = ["#", "Student", "Overall", "Turned in", "Score %", "Perf. index", "Rank", "Percentile", "z-score",
             "Curved", "Status", "", "tie"]
    for i, t in enumerate(heads):
        if t:
            hdr(sd.cell(row=7, column=i + 1), t, 9)
    sd.column_dimensions["M"].hidden = True
    idx = f"$F${ST_FIRST}:$F${ST_LAST}"
    for i in range(N_STU):
        r = ST_FIRST + i
        gr = GB_FIRST + i
        sd.cell(row=r, column=1, value=i + 1).font = ft(8, color=MUTED)
        calc(sd.cell(row=r, column=2), f'=Gradebook!B{gr}', bold=True)
        calc(sd.cell(row=r, column=3), f'=Gradebook!C{gr}', "0.0%", center=True)
        calc(sd.cell(row=r, column=4), f'=Gradebook!E{gr}', "0%", center=True)
        calc(sd.cell(row=r, column=5), f'=Gradebook!G{gr}', "0%", center=True)
        calc(sd.cell(row=r, column=6),
             f'=IF(B{r}="","",IF(AND(NOT(ISNUMBER(C{r})),NOT(ISNUMBER(D{r}))),"",'
             f'(IF(ISNUMBER(C{r}),Setup!$B$22*C{r},0)+IF(ISNUMBER(D{r}),Setup!$B$23*D{r},0))/'
             f'(IF(ISNUMBER(C{r}),Setup!$B$22,0)+IF(ISNUMBER(D{r}),Setup!$B$23,0))))', "0.0%", True, True)
        calc(sd.cell(row=r, column=7), f'=IF(ISNUMBER(F{r}),RANK(F{r},{idx},0),"")', "0", center=True)
        calc(sd.cell(row=r, column=8),
             f'=IF(ISNUMBER(F{r}),IF(COUNT({idx})<=1,1,SUMPRODUCT(--({idx}<F{r}))/(COUNT({idx})-1)),"")', "0%", center=True)
        calc(sd.cell(row=r, column=9), f'=IF(ISNUMBER(F{r}),IFERROR((F{r}-AVERAGE({idx}))/STDEVP({idx}),0),"")', "0.00", center=True)
        calc(sd.cell(row=r, column=10), f'=IF(ISNUMBER(H{r}),VLOOKUP(H{r},Setup!$G$12:$H$16,2,TRUE),"")', center=True)
        calc(sd.cell(row=r, column=11),
             f'=IF(B{r}="","",IF(OR(AND(ISNUMBER(C{r}),C{r}<Setup!$B$27),AND(ISNUMBER(D{r}),D{r}<Setup!$B$28)),"At risk",'
             f'IF(OR(AND(ISNUMBER(C{r}),C{r}<Setup!$B$27+0.1),AND(ISNUMBER(D{r}),D{r}<Setup!$B$28+0.1)),"Watch","On track")))',
             center=True)
        sd.cell(row=r, column=13, value=f'=IF(ISNUMBER(F{r}),G{r}+SUMPRODUCT(--($F$7:F{r - 1 if r > ST_FIRST else 7}=F{r})),"")')
    sd.conditional_formatting.add(f"K{ST_FIRST}:K{ST_LAST}", FormulaRule(formula=[f'K{ST_FIRST}="At risk"'],
                                                                         fill=PatternFill("solid", fgColor="F6D5D0"), font=Font(color="9A2A1C", bold=True)))
    sd.conditional_formatting.add(f"K{ST_FIRST}:K{ST_LAST}", FormulaRule(formula=[f'K{ST_FIRST}="Watch"'],
                                                                         fill=PatternFill("solid", fgColor="FCEBC9"), font=Font(color="7A5200", bold=True)))
    sd.conditional_formatting.add(f"F{ST_FIRST}:F{ST_LAST}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                                                         color="7FB0BF"))
    # ranked list
    hdr(sd["N7"], "Rank", 9)
    hdr(sd["O7"], "Student (highest first)", 9)
    hdr(sd["P7"], "Perf. index", 9)
    hdr(sd["Q7"], "Percentile", 9)
    hdr(sd["R7"], "Status", 9)
    tie = f"$M${ST_FIRST}:$M${ST_LAST}"
    for k in range(N_STU):
        r = ST_FIRST + k
        calc(sd.cell(row=r, column=14), f'=IF({k + 1}>COUNT({idx}),"",{k + 1})', "0", center=True)
        calc(sd.cell(row=r, column=15), f'=IF(N{r}="","",INDEX($B${ST_FIRST}:$B${ST_LAST},MATCH(N{r},{tie},0)))')
        calc(sd.cell(row=r, column=16), f'=IF(N{r}="","",INDEX({idx},MATCH(N{r},{tie},0)))', "0.0%", center=True)
        calc(sd.cell(row=r, column=17), f'=IF(N{r}="","",INDEX($H${ST_FIRST}:$H${ST_LAST},MATCH(N{r},{tie},0)))', "0%", center=True)
        calc(sd.cell(row=r, column=18), f'=IF(N{r}="","",INDEX($K${ST_FIRST}:$K${ST_LAST},MATCH(N{r},{tie},0)))', center=True)
    sd.conditional_formatting.add(f"R{ST_FIRST}:R{ST_LAST}", FormulaRule(formula=[f'R{ST_FIRST}="At risk"'],
                                                                         font=Font(color="9A2A1C", bold=True)))
    sd.freeze_panes = "C8"

    # ====================== Dashboard
    d = wb.create_sheet("Dashboard", 1)
    d.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKLMN", [2, 30, 16, 2, 30, 16, 2, 24, 10, 3, 6, 6, 6, 6]):
        d.column_dimensions[col].width = w
    banner(d, "Dashboard", "", 10)
    d["A2"] = '=Setup!B4&"  |  "&Setup!B5&"  |  as of "&TEXT(Setup!B8,"mmm d, yyyy")&"  |  week "&Pacing!C3&" of "&Setup!B7'
    d["A2"].font = ft(10, color="DCEBEF")
    kpis = [("B", "Class average", f'=IFERROR(AVERAGE(Gradebook!C{GB_FIRST}:C{GB_LAST}),"")', "0.0%"),
            ("E", "Work turned in", f'=IFERROR(AVERAGE(Gradebook!E{GB_FIRST}:E{GB_LAST}),"")', "0%"),
            ("H", "Students at risk", f'=COUNTIF(Standings!K{ST_FIRST}:K{ST_LAST},"At risk")', "0")]
    for col, lab, formula, fmt in kpis:
        c1 = d[f"{col}4"]
        c1.value = lab
        c1.font = ft(10, True, MUTED)
        c2 = d[f"{col}5"]
        c2.value = formula
        c2.number_format = fmt
        c2.font = ft(26, True, TEAL if col != "H" else "9A2A1C")
        d.row_dimensions[5].height = 38
        for rr in (4, 5, 6):
            d[f"{col}{rr}"].fill = PatternFill("solid", fgColor=TEAL_PALE)
            nxt = chr(ord(col) + 1)
            d[f"{nxt}{rr}"].fill = PatternFill("solid", fgColor=TEAL_PALE)
    d["B6"] = f'=COUNT(Gradebook!C{GB_FIRST}:C{GB_LAST})&" students graded"'
    d["E6"] = f'="Missing items: "&SUM(Gradebook!F{GB_FIRST}:F{GB_LAST})'
    d["H6"] = f'=COUNTIF(Standings!K{ST_FIRST}:K{ST_LAST},"Watch")&" more to watch"'
    for c in ("B6", "E6", "H6"):
        d[c].font = ft(9, color=MUTED)

    d["B8"] = "Grade spread"
    d["B8"].font = ft(11, True)
    for i, gl in enumerate(["A", "B", "C", "D", "F"]):
        r = 9 + i
        d.cell(row=r, column=11, value=gl).font = ft(8, color=MUTED)
        d.cell(row=r, column=12, value=f'=COUNTIF(Gradebook!D{GB_FIRST}:D{GB_LAST},"{gl}*")').font = ft(8, color=MUTED)
    d["K8"] = "Grade"
    d["L8"] = "Students"
    for c in ("K8", "L8"):
        d[c].font = ft(8, True, MUTED)
    bc = BarChart()
    bc.type = "col"
    bc.legend = None
    bc.title = None
    bc.y_axis.majorGridlines = None
    bc.add_data(Reference(d, min_col=12, min_row=8, max_row=13), titles_from_data=True)
    bc.set_categories(Reference(d, min_col=11, min_row=9, max_row=13))
    bc.series[0].graphicalProperties.solidFill = TEAL
    bc.height, bc.width = 6.4, 8.6
    d.add_chart(bc, "B9")

    d["E8"] = "Turned in vs. score"
    d["E8"].font = ft(11, True)
    sc_ch = ScatterChart()
    sc_ch.title = None
    sc_ch.style = 13
    sc_ch.legend = None
    sc_ch.x_axis.title = "Work turned in"
    sc_ch.y_axis.title = "Score % on turned-in work"
    sc_ch.x_axis.number_format = "0%"
    sc_ch.y_axis.number_format = "0%"
    sc_ch.x_axis.scaling.min, sc_ch.x_axis.scaling.max = 0, 1
    sc_ch.y_axis.scaling.min, sc_ch.y_axis.scaling.max = 0, 1
    xs = Reference(sd, min_col=4, min_row=ST_FIRST, max_row=ST_LAST)
    ys = Reference(sd, min_col=5, min_row=ST_FIRST, max_row=ST_LAST)
    ser = Series(ys, xs, title="Students")
    ser.marker.symbol = "circle"
    ser.marker.size = 7
    ser.marker.graphicalProperties.solidFill = ACCENT
    ser.marker.graphicalProperties.line.solidFill = ACCENT
    ser.graphicalProperties.line.noFill = True
    sc_ch.series.append(ser)
    sc_ch.height, sc_ch.width = 6.4, 8.6
    d.add_chart(sc_ch, "E9")

    d["H8"] = "Needs attention (lowest index)"
    d["H8"].font = ft(11, True)
    hdr(d["H9"], "Student", 9)
    hdr(d["I9"], "Index", 9)
    for k in range(8):
        r = 10 + k
        pos = f"(COUNT(Standings!$F${ST_FIRST}:$F${ST_LAST})-{k})"
        calc(d.cell(row=r, column=8),
             f'=IF({pos}<1,"",INDEX(Standings!$B${ST_FIRST}:$B${ST_LAST},MATCH({pos},Standings!$M${ST_FIRST}:$M${ST_LAST},0)))', size=9)
        calc(d.cell(row=r, column=9),
             f'=IF({pos}<1,"",INDEX(Standings!$F${ST_FIRST}:$F${ST_LAST},MATCH({pos},Standings!$M${ST_FIRST}:$M${ST_LAST},0)))',
             "0%", center=True, size=9)
    d["K20"] = "=Setup!B27"
    d["K20"].font = ft(7, color="FFFFFF")
    d.conditional_formatting.add("I10:I17", FormulaRule(formula=['AND(ISNUMBER(I10),I10<$K$20)'],
                                                        font=Font(color="9A2A1C", bold=True)))

    d["B23"] = "Workload by week (assignments due)"
    d["B23"].font = ft(11, True)
    wl = BarChart()
    wl.type = "col"
    wl.legend = None
    wl.y_axis.majorGridlines = None
    wl.add_data(Reference(pc, min_col=WK1, max_col=WKN, min_row=P_LAST + 2), from_rows=True, titles_from_data=False)
    wl.set_categories(Reference(pc, min_col=WK1, max_col=WKN, min_row=4))
    wl.series[0].graphicalProperties.solidFill = "3F7F92"
    wl.height, wl.width = 5.5, 17.6
    d.add_chart(wl, "B24")

    # ====================== Student Report
    sr = wb.create_sheet("Student Report")
    sr.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [4, 30, 16, 12, 8, 8, 14, 12]):
        sr.column_dimensions[col].width = w
    banner(sr, "Student Report", "", 8)
    sr["A2"] = '=Setup!B4&"  |  "&Setup!B5&"  |  as of "&TEXT(Setup!B8,"mmm d, yyyy")'
    sr["A2"].font = ft(10, color="DCEBEF")
    note(sr["B3"], "Select a student:")
    sr["B4"] = names[0] if names else ""
    sr["B4"].font = ft(13, True)
    sr["B4"].fill = PatternFill("solid", fgColor="FFF3C4")
    sr["B4"].border = BOX
    note(sr["C4"], "< pick a student")
    dv_s = DataValidation(type="list", formula1=f"=Roster!$B${R_FIRST}:$B${R_LAST}", allow_blank=True)
    sr.add_data_validation(dv_s)
    dv_s.add("B4")
    sr["H4"] = f'=IFERROR(MATCH(B4,Gradebook!$B${GB_FIRST}:$B${GB_LAST},0),"")'
    sr["H4"].font = ft(8, color="FFFFFF")
    k = "$H$4"
    metrics = [("Overall grade", f'=IF({k}="","",INDEX(Gradebook!$C${GB_FIRST}:$C${GB_LAST},{k}))', "0.0%"),
               ("Letter grade", f'=IF({k}="","",INDEX(Gradebook!$D${GB_FIRST}:$D${GB_LAST},{k}))', "@"),
               ("Work turned in", f'=IF({k}="","",INDEX(Gradebook!$E${GB_FIRST}:$E${GB_LAST},{k}))', "0%"),
               ("Missing items", f'=IF({k}="","",INDEX(Gradebook!$F${GB_FIRST}:$F${GB_LAST},{k}))', "0"),
               ("Class standing", f'=IF({k}="","",IF(ISNUMBER(INDEX(Standings!$G${ST_FIRST}:$G${ST_LAST},{k})),"#"&INDEX(Standings!$G${ST_FIRST}:$G${ST_LAST},{k})&" of "&COUNT(Standings!$F${ST_FIRST}:$F${ST_LAST})&"  (percentile "&TEXT(INDEX(Standings!$H${ST_FIRST}:$H${ST_LAST},{k}),"0%")&")",""))', "@"),
               ("Status", f'=IF({k}="","",INDEX(Standings!$K${ST_FIRST}:$K${ST_LAST},{k}))', "@")]
    for i, (lab, formula, fmt) in enumerate(metrics):
        r = 6 + i
        sr.cell(row=r, column=2, value=lab).font = ft(10, True)
        calc(sr.cell(row=r, column=3), formula, fmt, True)
        sr.cell(row=r, column=3).alignment = Alignment(horizontal="left", indent=1)
        sr.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    hdr(sr["G6"], "Category", 9)
    hdr(sr["H6"], "Score", 9)
    for j in range(5):
        r = 7 + j
        calc(sr.cell(row=r, column=7), f"=Setup!A{12 + j}", size=9)
        calc(sr.cell(row=r, column=8), f'=IF({k}="","",INDEX(Gradebook!${L(8 + j)}${GB_FIRST}:${L(8 + j)}${GB_LAST},{k}))',
             "0%", center=True, size=9)
    heads = ["#", "Assignment", "Unit", "Due", "Points", "Score", "%", "Status"]
    for i, t in enumerate(heads):
        hdr(sr.cell(row=14, column=i + 1), t, 9)
    for i in range(N_ASG):
        r = 15 + i
        ar = A_FIRST + i
        col = L(SC1 + i)
        sr.cell(row=r, column=1, value=f'=IF(Assignments!B{ar}="","",{i + 1})')
        sr.cell(row=r, column=2, value=f'=IF(A{r}="","",Assignments!B{ar})')
        sr.cell(row=r, column=3, value=f'=IF(A{r}="","",Assignments!C{ar})')
        sr.cell(row=r, column=4, value=f'=IF(A{r}="","",Assignments!G{ar})')
        sr.cell(row=r, column=4).number_format = "m/d"
        sr.cell(row=r, column=5, value=f'=IF(A{r}="","",Assignments!E{ar})')
        sr.cell(row=r, column=6, value=f'=IF(OR(A{r}="",{k}=""),"",IF(INDEX(Gradebook!{col}${GB_FIRST}:{col}${GB_LAST},{k})="","",INDEX(Gradebook!{col}${GB_FIRST}:{col}${GB_LAST},{k})))')
        sr.cell(row=r, column=7, value=f'=IF(OR(NOT(ISNUMBER(F{r})),N(E{r})=0),"",F{r}/E{r})')
        sr.cell(row=r, column=7).number_format = "0%"
        sr.cell(row=r, column=8, value=f'=IF(OR(A{r}="",{k}=""),"",IF(ISNUMBER(F{r}),"Turned in",IF(UPPER(F{r})="EX","Excused",'
                                          f'IF(OR(UPPER(F{r})="M",Assignments!I{ar}=1),"Missing","Not due yet"))))')
        for c in range(1, 9):
            cell = sr.cell(row=r, column=c)
            cell.font = ft(9)
            cell.border = BOX
            if c != 2 and c != 3:
                cell.alignment = Alignment(horizontal="center")
    sr.conditional_formatting.add(f"H15:H{14 + N_ASG}", FormulaRule(formula=['H15="Missing"'],
                                                                    fill=PatternFill("solid", fgColor="FBE0C6"), font=Font(color="8A3B00", bold=True)))
    sr.conditional_formatting.add(f"H15:H{14 + N_ASG}", FormulaRule(formula=['H15="Not due yet"'], font=Font(color="8A949B", italic=True)))
    sr.conditional_formatting.add("C11", FormulaRule(formula=['C11="At risk"'], font=Font(color="9A2A1C", bold=True)))
    sr.print_area = f"A1:H{14 + N_ASG}"
    sr.page_setup.fitToWidth = 1
    sr.page_setup.fitToHeight = 0
    sr.sheet_properties.pageSetUpPr.fitToPage = True

    tabs = {"Start Here": ACCENT, "Dashboard": ACCENT, "Setup": TEAL, "Roster": TEAL, "Pacing": TEAL, "Assignments": TEAL,
            "Gradebook": TEAL, "Standings": "5E6C75", "Student Report": "5E6C75"}
    for w in wb.worksheets:
        w.sheet_properties.tabColor = tabs[w.title]
    # order: Start Here, Dashboard, Setup, Roster, Pacing, Assignments, Gradebook, Standings, Student Report
    wb._sheets = [wb[n] for n in ["Start Here", "Dashboard", "Setup", "Roster", "Pacing", "Assignments", "Gradebook",
                                  "Standings", "Student Report"]]
    wb.active = 0
    wb.save(path)


if __name__ == "__main__":
    out = sys.argv[1]
    mode = sys.argv[2] if len(sys.argv) > 2 else "all"
    ts, units, asg, names, profile, rnd = sample()
    asof = date(2026, 10, 9)
    scores = fill_scores(ts, units, asg, names, profile, rnd, asof)
    # sample for buyers: live date. test/screenshot copy: fixed date.
    build(f"{out}/Teacher_Command_Center_Sample.xlsx", ts, units, asg, names, scores, asof_value="=TODAY()")
    build(f"{out}/cc_test_fixed_date.xlsx", ts, units, asg, names, scores, asof_value=asof)
    build(f"{out}/Teacher_Command_Center_Blank.xlsx", ts, [("Unit 1", 1, 3)],
          [("Unit 1: Homework 1", "Unit 1", "Homework", 10, None)], ["Example Student"], {(0, 0): 9},
          asof_value="=TODAY()", class_name="My Class", teacher="")
    print("built")
