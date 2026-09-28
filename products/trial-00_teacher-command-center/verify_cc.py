"""Independent re-computation of the Teacher Command Center, compared with LibreOffice-recalculated values."""
import math
import sys
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

GB_FIRST, N_STU, N_ASG, SC1 = 10, 40, 60, 13
A_FIRST, P_FIRST, N_UNIT, ST_FIRST = 6, 6, 20, 8

wb = load_workbook(sys.argv[1], data_only=True)
S, P, A, G, SD, D, R = (wb[n] for n in ["Setup", "Pacing", "Assignments", "Gradebook", "Standings", "Dashboard", "Student Report"])


def d(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(v))
    return None


term = d(S["B6"].value)
asof = d(S["B8"].value)
blank_zero = S["B9"].value == "Yes"
cats = [(S.cell(row=12 + i, column=1).value, S.cell(row=12 + i, column=2).value) for i in range(5)]
scale = sorted((S.cell(row=12 + i, column=4).value, S.cell(row=12 + i, column=5).value) for i in range(12))
curve = sorted((S.cell(row=12 + i, column=7).value, S.cell(row=12 + i, column=8).value) for i in range(5))
w_over, w_sub = S["B22"].value, S["B23"].value
th_over, th_sub = S["B27"].value, S["B28"].value

unit_end = {}
for i in range(N_UNIT):
    r = P_FIRST + i
    name, sw, wk = P.cell(row=r, column=2).value, P.cell(row=r, column=3).value, P.cell(row=r, column=4).value
    if name and sw:
        unit_end[name] = term + timedelta(days=(sw + max(1, wk or 0) - 2) * 7 + 4)

asg = []
for i in range(N_ASG):
    r = A_FIRST + i
    name = A.cell(row=r, column=2).value
    if not name:
        asg.append(None)
        continue
    unit, cat, pts, manual = (A.cell(row=r, column=c).value for c in (3, 4, 5, 6))
    due = d(manual) or unit_end.get(unit)
    asg.append(dict(name=name, unit=unit, cat=cat, pts=pts or 0, due=due, isdue=1 if (due and due <= asof) else 0))


def letter(p, table):
    out = table[0][1]
    for mn, l in table:
        if p >= mn - 1e-12:
            out = l
    return out


problems = []


def close(a, b, tol=1e-9):
    if a is None or b is None or a == "" or b == "":
        return (a in (None, "")) and (b in (None, ""))
    if isinstance(a, str) or isinstance(b, str):
        return a == b
    return abs(a - b) <= tol


def chk(label, exp, got):
    if not close(exp, got):
        problems.append((label, exp, got))


students = []
for s in range(N_STU):
    r = GB_FIRST + s
    name = G.cell(row=r, column=2).value
    if not name:
        continue
    sc = [G.cell(row=r, column=SC1 + i).value for i in range(N_ASG)]
    isnum = [isinstance(v, (int, float)) for v in sc]
    ism = [isinstance(v, str) and v.upper() == "M" for v in sc]
    isex = [isinstance(v, str) and v.upper() == "EX" for v in sc]
    isblank = [v is None or v == "" for v in sc]
    pct = {}
    for cat, _ in cats:
        earned = sum(v for a, v, n in zip(asg, sc, isnum) if a and a["cat"] == cat and n)
        possible = sum(a["pts"] * (n + m + (blank_zero * a["isdue"] * b))
                       for a, n, m, b in zip(asg, isnum, ism, isblank) if a and a["cat"] == cat)
        pct[cat] = earned / possible if possible else None
    used = [(w, pct[c]) for c, w in cats if pct[c] is not None]
    den = sum(w for w, _ in used)
    overall = sum(w * p for w, p in used) / den if used and den else None
    sub_den = sum(a["isdue"] * (not e) for a, e in zip(asg, isex) if a)
    sub_num = sum(a["isdue"] * n for a, n in zip(asg, isnum) if a)
    submit = sub_num / sub_den if sub_den else None
    missing = sum(a["isdue"] * b for a, b in zip(asg, isblank) if a) + sum(ism)
    pts_sub = sum(a["pts"] * n for a, n in zip(asg, isnum) if a)
    scorepct = sum(v for v, n in zip(sc, isnum) if n) / pts_sub if pts_sub else None
    chk(f"{name} overall", overall, G.cell(row=r, column=3).value)
    chk(f"{name} letter", letter(overall, scale) if overall is not None else None, G.cell(row=r, column=4).value)
    chk(f"{name} submit", submit, G.cell(row=r, column=5).value)
    chk(f"{name} missing", missing, G.cell(row=r, column=6).value)
    chk(f"{name} score%", scorepct, G.cell(row=r, column=7).value)
    for j, (cat, _) in enumerate(cats):
        chk(f"{name} {cat}", pct[cat], G.cell(row=r, column=8 + j).value)
    if overall is None and submit is None:
        idx = None
    else:
        idx = ((w_over * overall if overall is not None else 0) + (w_sub * submit if submit is not None else 0)) / \
              ((w_over if overall is not None else 0) + (w_sub if submit is not None else 0))
        idx = round(idx, 9)  # the workbook rounds the index to 9 decimals so float noise cannot split ties
    students.append(dict(row=s, name=name, overall=overall, submit=submit, idx=idx, missing=missing))

vals = [x["idx"] for x in students if x["idx"] is not None]
n = len(vals)
mean = sum(vals) / n if n else 0
sd = math.sqrt(sum((v - mean) ** 2 for v in vals) / n) if n else 0
for x in students:
    r = ST_FIRST + x["row"]
    chk(f"{x['name']} index", x["idx"], SD.cell(row=r, column=6).value)
    if x["idx"] is None:
        continue
    rank = 1 + sum(1 for v in vals if v > x["idx"] + 1e-12)
    pctile = 1 if n <= 1 else sum(1 for v in vals if v < x["idx"] - 1e-12) / (n - 1)
    z = (x["idx"] - mean) / sd if sd else 0
    status = "At risk" if ((x["overall"] is not None and x["overall"] < th_over) or (x["submit"] is not None and x["submit"] < th_sub)) \
        else ("Watch" if ((x["overall"] is not None and x["overall"] < th_over + 0.1) or (x["submit"] is not None and x["submit"] < th_sub + 0.1))
              else "On track")
    x.update(rank=rank, pctile=pctile, status=status)
    chk(f"{x['name']} rank", rank, SD.cell(row=r, column=7).value)
    chk(f"{x['name']} percentile", pctile, SD.cell(row=r, column=8).value)
    chk(f"{x['name']} z", z, SD.cell(row=r, column=9).value)
    chk(f"{x['name']} curved", letter(pctile, curve), SD.cell(row=r, column=10).value)
    chk(f"{x['name']} status", status, SD.cell(row=r, column=11).value)

ranked = sorted([x for x in students if x["idx"] is not None], key=lambda x: (x["rank"], x["row"]))
for k, x in enumerate(ranked):
    chk(f"ranked #{k + 1} name", x["name"], SD.cell(row=ST_FIRST + k, column=15).value)
low = list(reversed(ranked))[:8]
for k, x in enumerate(low):
    chk(f"attention #{k + 1}", x["name"], D.cell(row=10 + k, column=8).value)

graded = [x["overall"] for x in students if x["overall"] is not None]
chk("dash class avg", sum(graded) / len(graded), D["B5"].value)
subs = [x["submit"] for x in students if x["submit"] is not None]
chk("dash submit avg", sum(subs) / len(subs), D["E5"].value)
chk("dash at risk", sum(1 for x in students if x.get("status") == "At risk"), D["H5"].value)

for i, a in enumerate(asg):
    if not a:
        continue
    r = A_FIRST + i
    col = [G.cell(row=GB_FIRST + x["row"], column=SC1 + i).value for x in students]
    nums = [v for v in col if isinstance(v, (int, float))]
    exc = sum(1 for v in col if isinstance(v, str) and v.upper() == "EX")
    chk(f"asg {i + 1} week", (a["due"] - term).days // 7 + 1 if a["due"] else None, A.cell(row=r, column=8).value)
    chk(f"asg {i + 1} due?", a["isdue"], A.cell(row=r, column=9).value)
    chk(f"asg {i + 1} class avg", (sum(nums) / len(nums) / a["pts"]) if nums and a["pts"] else None, A.cell(row=r, column=10).value)
    if a["isdue"]:
        chk(f"asg {i + 1} turned in", len(nums) / (len(students) - exc) if len(students) - exc else None, A.cell(row=r, column=11).value)
    m = sum(1 for v in col if isinstance(v, str) and v.upper() == "M") + (sum(1 for v in col if v is None) if a["isdue"] else 0)
    chk(f"asg {i + 1} missing", m, A.cell(row=r, column=12).value)

for w in range(20):
    exp = sum(1 for a in asg if a and a["due"] and (a["due"] - term).days // 7 + 1 == w + 1)
    chk(f"week {w + 1} load", exp, P.cell(row=27, column=8 + w).value)

sel = R["B4"].value
x = next((s for s in students if s["name"] == sel), None)
if x:
    for i, a in enumerate(asg):
        if not a:
            continue
        v = G.cell(row=GB_FIRST + x["row"], column=SC1 + i).value
        exp = "Submitted" if isinstance(v, (int, float)) else ("Excused" if isinstance(v, str) and v.upper() == "EX" else
                                                                ("Missing" if (isinstance(v, str) and v.upper() == "M") or a["isdue"] else "Not due yet"))
        chk(f"report {i + 1}", exp, R.cell(row=15 + i, column=8).value)

print(f"students: {len(students)}  assignments: {sum(1 for a in asg if a)}  due: {sum(a['isdue'] for a in asg if a)}  problems: {len(problems)}")
for p in problems[:25]:
    print(p)
