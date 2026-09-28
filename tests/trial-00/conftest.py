"""Fixtures for TRIAL-00 acceptance tests (R6, design v1, ctx-r6-trial00-20260927)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from recalc import recalc  # noqa: E402

REPO = HERE.parents[1]
PRODUCT = REPO / "products" / "trial-00_teacher-command-center"
SAMPLE = PRODUCT / "Teacher_Command_Center_Sample.xlsx"
BLANK = PRODUCT / "Teacher_Command_Center_Blank.xlsx"
LISTING_DIR = PRODUCT / "listing"
LISTING_JSON = LISTING_DIR / "listing.json"

DESIGN_VERSION = "v1"
CONTEXT_ID = "ctx-r6-trial00-20260927"


@pytest.fixture(scope="session")
def workbooks():
    out = {}
    for label, p in (("Sample", SAMPLE), ("Blank", BLANK)):
        if not p.exists():
            pytest.fail(f"{label} workbook missing: {p}")
        out[label] = openpyxl.load_workbook(p)  # formulas, not values
    return out


@pytest.fixture(scope="session")
def lo_tmp(tmp_path_factory):
    return tmp_path_factory.mktemp("recalc")


_LO_STATE: dict = {}


def lo_available(tmpdir: Path) -> tuple[bool, str]:
    """Probe whether LibreOffice can actually load and recalc an xlsx here."""
    if "ok" in _LO_STATE:
        return _LO_STATE["ok"], _LO_STATE["why"]
    probe = tmpdir / "probe.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 2
    wb.active["A2"] = "=A1*21"
    wb.save(probe)
    try:
        out = recalc(probe, tmpdir / "probe_out", timeout=120)
        v = openpyxl.load_workbook(out, data_only=True).active["A2"].value
        ok = v == 42
        why = "" if ok else f"probe recalculated to {v!r} instead of 42"
    except Exception as e:  # noqa: BLE001
        ok, why = False, str(e).splitlines()[0] + " | " + " ".join(str(e).split())[-200:]
    _LO_STATE.update(ok=ok, why=why)
    return ok, why


@pytest.fixture(scope="session")
def require_lo(lo_tmp):
    ok, why = lo_available(lo_tmp)
    if not ok:
        pytest.skip(
            "BLOCKED: LibreOffice Calc cannot load/recalculate xlsx in this environment "
            f"({why}). 未実行 — 合格ではない。T00-ENV-01 を参照。"
        )
    return True


@pytest.fixture(scope="session")
def recalculated(require_lo, lo_tmp):
    """Cache: label -> path of a LibreOffice-recalculated copy."""
    cache: dict = {}

    def get(label: str, src: Path, run: int = 1) -> Path:
        key = (label, run)
        if key not in cache:
            cache[key] = recalc(src, lo_tmp / f"{label}_{run}")
        return cache[key]

    return get


@pytest.fixture(scope="session")
def listing():
    """Listing package contract. A missing listing is a real T4 defect: tests call need_listing(), which
    FAILS inside the test body (not skip, not a fixture error)."""
    if not LISTING_JSON.exists():
        return {"__error__": f"listing package missing: {LISTING_JSON.relative_to(REPO)} does not exist (T4 blocker)"}
    try:
        data = json.loads(LISTING_JSON.read_text(encoding="utf-8"))
    except Exception as e:  # noqa: BLE001
        return {"__error__": f"listing.json is not valid JSON: {e}"}
    if not isinstance(data, dict):
        return {"__error__": "listing.json top level is not an object"}
    return data


def need_listing(listing: dict) -> dict:
    if "__error__" in listing:
        pytest.fail(listing["__error__"])
    return listing


def resolve_listing_path(rel: str) -> Path | None:
    for base in (LISTING_DIR, PRODUCT):
        p = (base / rel).resolve()
        if p.exists():
            return p
    return None


@pytest.fixture(scope="session")
def description_text(listing):
    if "__error__" in listing:
        return None
    rel = listing.get("description_file")
    p = resolve_listing_path(rel) if isinstance(rel, str) else None
    return p.read_text(encoding="utf-8") if p is not None else None


def need_text(listing: dict, text):
    need_listing(listing)
    if text is None:
        pytest.fail(f"description_file {listing.get('description_file')!r} not found")
    return text
