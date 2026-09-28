"""Build input scenarios on a COPY of the Blank workbook (R6, ctx-r6-trial00-20260927).

Cell locations (input areas) were located from header labels in the Blank file:
Setup B8 (As of), B9 (count past-due blanks), Roster B6:B45, Pacing B6:D25,
Assignments B6:F65, Gradebook grid M10:BT49 (row 4 = assignment number),
Student Report B4 (selected student). Each location is asserted against its
header before writing, so a layout change makes the harness fail loudly.
"""
from __future__ import annotations

import datetime as dt
import random
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

CATS = ["Homework", "Quizzes", "Tests", "Projects", "Participation"]
TERM_START = dt.date(2026, 8, 17)  # Monday, as in the Blank Setup


def _assert(cond, msg):
    if not cond:
        raise AssertionError(f"HARNESS layout check failed: {msg}")


def _clear(ws, rows, cols):
    for r in rows:
        for c in cols:
            ws.cell(r, c).value = None


def build(blank: Path, out: Path, *, as_of: dt.date, count_blanks: str = "Yes",
          units=None, assignments=None, students=None, scores=None) -> Path:
    """assignments: list of (name, unit, category, points, due_date_or_None)
    scores: dict[(student_idx, assignment_no)] -> value (number, "M", "EX", None)"""
    wb = openpyxl.load_workbook(blank)
    su, ro, pa, asg, gb, rp = (wb[n] for n in ["Setup", "Roster", "Pacing", "Assignments", "Gradebook", "Student Report"])
    _assert(str(su["A8"].value).startswith("As of"), "Setup A8")
    _assert(str(su["A9"].value).startswith("Count past-due"), "Setup A9")
    _assert(ro["B5"].value == "Student", "Roster B5")
    _assert(pa["B5"].value == "Unit" and pa["C5"].value == "Start wk", "Pacing B5:C5")
    _assert([asg.cell(5, c).value for c in range(2, 7)] == ["Assignment", "Unit", "Category", "Points", "Due date (optional)"], "Assignments B5:F5")
    _assert(gb["M4"].value == 1 and gb["BT4"].value == 60, "Gradebook M4/BT4 assignment numbers")
    _assert(str(rp["B3"].value).startswith("Select a student"), "Student Report B3")

    su["B8"].value = dt.datetime.combine(as_of, dt.time())
    su["B9"].value = count_blanks
    units = units if units is not None else [("Unit 1", 1, 3), ("Unit 2", 4, 3), ("Unit 3", 7, 3)]
    _clear(pa, range(6, 26), range(2, 5))
    for j, (u, sw, wk) in enumerate(units):
        pa.cell(6 + j, 2).value, pa.cell(6 + j, 3).value, pa.cell(6 + j, 4).value = u, sw, wk
    _clear(asg, range(6, 66), range(2, 7))
    for j, (name, unit, cat, pts, due) in enumerate(assignments or []):
        r = 6 + j
        asg.cell(r, 2).value = name
        asg.cell(r, 3).value = unit
        asg.cell(r, 4).value = cat
        asg.cell(r, 5).value = pts
        asg.cell(r, 6).value = dt.datetime.combine(due, dt.time()) if due else None
    _clear(ro, range(6, 46), range(2, 5))
    for j, n in enumerate(students or []):
        ro.cell(6 + j, 2).value = n
    _clear(gb, range(10, 50), range(13, 73))
    for (i, k), v in (scores or {}).items():
        gb.cell(10 + i, 12 + k).value = v
    rp["B4"].value = (students or [None])[0]
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def names(n):
    first = ["Alex", "Blair", "Casey", "Drew", "Emery", "Finley", "Gray", "Harper", "Indy", "Jules"]
    return [f"Test Student {first[j % 10]} {j + 1:02d}" for j in range(n)]


def gen_assignments(n, due=dt.date(2026, 9, 4), unit="Unit 1"):
    pts = [10, 20, 50, 100, 5]
    return [(f"Item {k + 1}", unit, CATS[k % 5], pts[k % 5], due + dt.timedelta(days=k)) for k in range(n)]


def gen_scores(n_students, asgs, seed, p_m=0.08, p_ex=0.05, p_blank=0.05):
    rnd = random.Random(seed)
    out = {}
    for i in range(n_students):
        for k, (_, _, _, pts, _) in enumerate(asgs, start=1):
            x = rnd.random()
            if x < p_m:
                v = "M"
            elif x < p_m + p_ex:
                v = "EX"
            elif x < p_m + p_ex + p_blank:
                v = None
            else:
                v = round(rnd.uniform(0.4, 1.0) * pts, 1)
            out[(i, k)] = v
    return out


def scenarios():
    """name -> kwargs for build(). Fixed, deterministic inputs (design v1 section 2.1/2.3/2.4)."""
    late = dt.date(2026, 12, 31)
    s = {}
    a6 = gen_assignments(6)
    s["FIT-05a_5x6"] = dict(as_of=late, assignments=a6, students=names(5), scores=gen_scores(5, a6, 1))
    a12 = gen_assignments(12)
    s["FIT-05b_20x12"] = dict(as_of=late, assignments=a12, students=names(20), scores=gen_scores(20, a12, 2))
    a10 = gen_assignments(10)
    s["FIT-05c_allcats"] = dict(as_of=late, assignments=a10, students=names(8),
                                scores=gen_scores(8, a10, 3, p_m=0, p_ex=0, p_blank=0))
    a3 = gen_assignments(3)
    sc = {(0, 1): 8, (0, 2): None, (0, 3): 40, (1, 1): None, (1, 2): 15, (1, 3): None, (2, 1): 10, (2, 2): 20, (2, 3): 50}
    s["EDGE-01_blank_No"] = dict(as_of=late, count_blanks="No", assignments=a3, students=names(3), scores=sc)
    s["EDGE-01b_blank_Yes"] = dict(as_of=late, count_blanks="Yes", assignments=a3, students=names(3), scores=sc)
    s["EDGE-02_missing"] = dict(as_of=late, assignments=a3, students=names(3),
                                scores={**{(i, k): 0.8 * a3[k - 1][3] for i in range(3) for k in (1, 2, 3)}, (1, 2): "M"})
    s["EDGE-03_excused"] = dict(as_of=late, assignments=a3, students=names(3),
                                scores={**{(i, k): 0.8 * a3[k - 1][3] for i in range(3) for k in (1, 2, 3)}, (1, 2): "EX"})
    a5 = gen_assignments(5)
    s["EDGE-04_all_tied"] = dict(as_of=late, assignments=a5, students=names(4),
                                 scores={(i, k): 0.75 * a5[k - 1][3] for i in range(4) for k in range(1, 6)})
    s["EDGE-05_one_student"] = dict(as_of=late, assignments=a5, students=names(1),
                                    scores={(0, k): 0.9 * a5[k - 1][3] for k in range(1, 6)})
    s["EDGE-06_nothing_due_all_blank"] = dict(as_of=dt.date(2026, 8, 18), assignments=a5, students=names(6), scores={})
    a60 = gen_assignments(60)
    s["OOS-02_capacity_40x60"] = dict(as_of=late, assignments=a60, students=names(40), scores=gen_scores(40, a60, 4))
    return s
