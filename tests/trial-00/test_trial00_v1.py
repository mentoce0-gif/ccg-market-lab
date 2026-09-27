"""TRIAL-00 "Teacher Command Center v2" acceptance tests.

Design: tests/trial-00/test_design_v1.md (design version v1)
Design context ID: ctx-r6-trial00-20260927   Date: 2026-09-27
後付け：商品は本テストより前に作られた（v2 §8.4）。
R6 did not read build_command_center.py, verify_cc.py, listing/, driver/.
Severity per v2 §7.5 is given in each test's docstring as [S0]..[S3].
Manual tests are skipped with reason "MANUAL: ..." and are 未実行 (never a pass).
"""
from __future__ import annotations

import hashlib
import json
import re
import struct
import zipfile
import zlib
from pathlib import Path

import openpyxl
import pytest

import cases
import recompute
from conftest import BLANK, PRODUCT, SAMPLE, lo_available, need_listing, need_text, resolve_listing_path
from recalc import is_error_value, recalc

EXPECTED_TABS = ["Start Here", "Dashboard", "Setup", "Roster", "Pacing",
                 "Assignments", "Gradebook", "Standings", "Student Report"]
CLAIMED_FORMULA_COUNT = 2344  # 付録A / TRIAL-00

# Excel 2007 allowlist (design v1 §2.2). STDEVP/VARP etc. added at implementation (see 実装時の注記).
ALLOWED_FUNCS = set("""
SUM SUMIF SUMIFS SUMPRODUCT COUNT COUNTA COUNTBLANK COUNTIF COUNTIFS AVERAGE AVERAGEIF AVERAGEIFS
MIN MAX LARGE SMALL RANK PERCENTRANK PERCENTILE MEDIAN ROUND ROUNDUP ROUNDDOWN INT ABS MOD IF IFERROR
AND OR NOT TRUE FALSE ISBLANK ISNUMBER ISTEXT ISERROR ISNA INDEX MATCH VLOOKUP HLOOKUP LOOKUP CHOOSE
OFFSET INDIRECT ROW ROWS COLUMN COLUMNS TEXT LEFT RIGHT MID LEN TRIM UPPER LOWER PROPER CONCATENATE
REPT SUBSTITUTE FIND SEARCH VALUE EXACT DATE TODAY NOW YEAR MONTH DAY WEEKDAY WEEKNUM NETWORKDAYS
WORKDAY EDATE EOMONTH DATEDIF N T NA MAXA MINA STDEV SUBTOTAL HYPERLINK ISODD ISEVEN CEILING FLOOR
STDEVP STDEVA STDEVPA VAR VARP
""".split())
FORBIDDEN_NEWER = {"IFNA", "XLOOKUP", "FILTER", "LET", "IFS", "MAXIFS", "MINIFS", "TEXTJOIN", "CONCAT",
                   "SWITCH", "UNIQUE", "SORT", "SORTBY", "SEQUENCE", "XMATCH", "AGGREGATE", "RANK.EQ",
                   "RANK.AVG", "PERCENTILE.INC", "PERCENTILE.EXC", "STDEV.S", "STDEV.P", "LAMBDA"}

CJK = re.compile(r"[぀-ヿ㐀-鿿＀-￯]")
LEFTOVER = re.compile(r"\bTODO\b|\bFIXME\b|lorem|\[未検証\]|\bXXX\b|\bTBD\b", re.I)
STR_LIT = re.compile(r'"(?:[^"]|"")*"')
SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_\.]*))!")
FUNC = re.compile(r"(?<![A-Za-z0-9_\.])([A-Za-z_][A-Za-z0-9_\.]*)\s*\(")


# ------------------------------------------------------------------ helpers
def formulas(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c.coordinate, c.value


def strip_strings(f: str) -> str:
    return STR_LIT.sub('""', f)


def refs_other_sheets(f: str, own: str) -> set[str]:
    out = set()
    for q, u in SHEET_REF.findall(strip_strings(f)):
        name = (q or u).replace("''", "'")
        if name != own:
            out.add(name)
    return out


def norm_tab(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\(.*?\)", "", s)).strip().lower()


def visible(wb):
    return [ws for ws in wb.worksheets if ws.sheet_state == "visible"]


def error_cells(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return [(ws.title, c.coordinate, c.value) for ws in wb.worksheets
            for row in ws.iter_rows() for c in row if is_error_value(c.value)]


def all_values(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if c.value is not None}


def app_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("docProps/app.xml").decode("utf-8", "replace")


def xml_field(xml: str, tag: str) -> str | None:
    m = re.search(rf"<(?:\w+:)?{tag}[^>]*>(.*?)</(?:\w+:)?{tag}>", xml, re.S)
    return m.group(1) if m else None


def start_here_text(wb) -> str:
    return "\n".join(str(c.value) for row in wb["Start Here"].iter_rows() for c in row if c.value is not None)


BOOKS = pytest.mark.parametrize("label", ["Sample", "Blank"])


# ================================================================== ENV
def test_T00_ENV_01_libreoffice_recalc_available(lo_tmp):
    """[環境] LibreOffice Calc must be able to recalculate xlsx, otherwise every §7.3 recalculation test is
    未実行. This is an environment defect, not a product defect, but T3 cannot pass while it fails."""
    ok, why = lo_available(lo_tmp)
    assert ok, f"LibreOffice cannot recalculate xlsx here: {why}"


def test_T00_HARN_01_recompute_selfcheck():
    """[harness] R6's independent recomputation on a hand-worked toy case (not a product test)."""
    import datetime as dt
    A = recompute.Assignment
    m = recompute.Model(
        as_of=dt.date(2026, 10, 1), count_blanks=True,
        weights={"Homework": 0.5, "Tests": 0.5, "Projects": 0.0},
        scale=[(0.0, "F"), (0.7, "C"), (0.9, "A")], perf_w_overall=0.7, perf_w_sub=0.3,
        risk_overall=0.7, risk_sub=0.8, students=["P", "Q", "R"],
        assignments={1: A(1, "h1", "Homework", 10, dt.date(2026, 9, 1)),
                     2: A(2, "h2", "Homework", 10, dt.date(2026, 9, 2)),
                     3: A(3, "t1", "Tests", 100, dt.date(2026, 9, 3)),
                     4: A(4, "t2", "Tests", 100, dt.date(2026, 12, 1))},  # not due
        grid={(0, 1): 10, (0, 2): "M", (0, 3): 90, (0, 4): None,
              (1, 1): "EX", (1, 2): None, (1, 3): 80, (1, 4): 100,
              (2, 1): 10, (2, 2): "M", (2, 3): 90, (2, 4): None})
    r = recompute.class_metrics(m)
    # P: HW 10/20=0.5, Tests 90/100=0.9 -> overall 0.7; sub: due {1,2,3} turned 2/3; missing 1
    assert abs(r[0]["pct"]["Homework"] - 0.5) < 1e-12 and abs(r[0]["overall"] - 0.7) < 1e-12
    assert abs(r[0]["sub"] - 2 / 3) < 1e-12 and r[0]["missing"] == 1 and r[0]["grade"] == "C"
    # Q: HW: EX skipped, blank past due -> 0/10 -> 0; Tests 180/200 = 0.9; overall 0.45
    assert abs(r[1]["pct"]["Homework"] - 0.0) < 1e-12 and abs(r[1]["pct"]["Tests"] - 0.9) < 1e-12
    assert abs(r[1]["overall"] - 0.45) < 1e-12 and abs(r[1]["sub"] - 0.5) < 1e-12 and r[1]["missing"] == 1
    assert r[1]["statuses"] == {1: "Excused", 2: "Missing", 3: "Submitted", 4: "Submitted"}
    # P and R identical -> same rank 1; Q rank 3
    assert r[0]["rank"] == 1 and r[2]["rank"] == 1 and r[1]["rank"] == 3
    # P's overall sits exactly on the 0.70 line -> flagged as boundary (status not compared); Q is At risk
    assert r[1]["status"] == "At risk" and r[0]["status_boundary"] is True
    assert r[0]["pct"]["Projects"] is None


# ================================================================== 2.1 FIT
def test_T00_FIT_01_sample_and_blank_present():
    """[S1] §7.3 both Sample and Blank are shipped and open."""
    for p in (SAMPLE, BLANK):
        assert p.exists(), f"missing {p.name}"
        openpyxl.load_workbook(p)


@BOOKS
def test_T00_FIT_02_nine_tabs(workbooks, label):
    """[S2] 付録A: exactly 9 visible tabs, names map 1:1 to the listed tabs."""
    names = [ws.title for ws in visible(workbooks[label])]
    assert len(names) == 9, names
    assert sorted(norm_tab(n) for n in names) == sorted(norm_tab(n) for n in EXPECTED_TABS), names


def test_T00_FIT_03_linkage_chain(workbooks):
    """[S2] 付録A linkage: Pacing -> Assignments -> Gradebook(Roster x) -> Standings -> Report/Dashboard."""
    refs: dict[str, set] = {}
    for sheet, _, f in formulas(workbooks["Sample"]):
        refs.setdefault(sheet, set()).update(refs_other_sheets(f, sheet))
    need = [("Assignments", {"Pacing"}), ("Gradebook", {"Roster"}), ("Gradebook", {"Assignments"}),
            ("Standings", {"Gradebook"}), ("Student Report", {"Gradebook", "Standings"})]
    missing = [(a, b) for a, b in need if not (refs.get(a, set()) & b)]
    if not refs.get("Dashboard"):
        missing.append(("Dashboard", {"any"}))
    assert not missing, f"missing links: {missing}; found: { {k: sorted(v) for k, v in refs.items()} }"


def test_T00_FIT_04_formula_count_matches_record(workbooks):
    """[S3] 付録A claims 2,344 formulas; the record must match the product."""
    n = sum(1 for _ in formulas(workbooks["Sample"]))
    assert n == CLAIMED_FORMULA_COUNT, f"Sample has {n} formula cells, record says {CLAIMED_FORMULA_COUNT}"


# ================================================================== 2.2 SS (recalc)
def test_T00_SS_01_sample_recalc_zero_errors(recalculated):
    """[S1] §7.3 LibreOffice recalculation of Sample has 0 error values."""
    errs = error_cells(recalculated("Sample", SAMPLE))
    assert not errs, f"{len(errs)} error cells, e.g. {errs[:10]}"


def test_T00_SS_02_blank_recalc_zero_errors(recalculated):
    """[S1] §7.3 Blank before input shows no errors."""
    errs = error_cells(recalculated("Blank", BLANK))
    assert not errs, f"{len(errs)} error cells, e.g. {errs[:10]}"


@pytest.mark.skip(reason="MANUAL: T00-SS-03 Blank visual check before input (0%/broken borders) — 未実行")
def test_T00_SS_03_blank_visual_manual():
    """[S3]"""


def test_T00_SS_04_independent_recompute_sample(recalculated):
    """[S1] §7.3 independent Python recomputation of Sample: 0 mismatches."""
    wb = openpyxl.load_workbook(recalculated("Sample", SAMPLE), data_only=True)
    mism = recompute.compare_workbook(wb) + recompute.pacing_end_checks(wb)
    assert not mism, f"{len(mism)} mismatches, e.g. {mism[:15]}"


@BOOKS
def test_T00_SS_05_excel2007_functions_only(workbooks, label):
    """[S2] §7.3 only Excel-2007-era functions; no _xlfn. prefix."""
    bad = {}
    for sheet, coord, f in formulas(workbooks[label]):
        s = strip_strings(f)
        if "_xlfn" in s.lower() or "_xlws" in s.lower():
            bad.setdefault("_xlfn", []).append(f"{sheet}!{coord}")
        for fn in FUNC.findall(s):
            u = fn.upper()
            if u not in ALLOWED_FUNCS:
                bad.setdefault(u, []).append(f"{sheet}!{coord}")
    assert not bad, {k: (len(v), v[:3]) for k, v in bad.items()}


@BOOKS
def test_T00_SS_06_conditional_formatting_same_sheet(workbooks, label):
    """[S2] §7.3 conditional formatting references only its own sheet (Google Sheets limitation)."""
    wb = workbooks[label]
    names = set(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else set()
    bad = []
    for ws in wb.worksheets:
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                texts = list(rule.formula or [])
                for sub in (getattr(rule, "dataBar", None), getattr(rule, "colorScale", None), getattr(rule, "iconSet", None)):
                    for v in (getattr(sub, "cfvo", None) or []):
                        if v.val is not None:
                            texts.append(str(v.val))
                for t in texts:
                    s = strip_strings(t)
                    if "!" in s or "[" in s or any(re.search(rf"\b{re.escape(n)}\b", s) for n in names):
                        bad.append((ws.title, str(cf.sqref), t))
    assert not bad, bad


@pytest.mark.skip(reason="MANUAL: T00-SS-07 Google Sheets parity (open, compare overall/rank with Excel) — 未実行 (§17 #2)")
def test_T00_SS_07_google_sheets_parity_manual():
    """[S1]"""


@BOOKS
def test_T00_SS_08_delivered_is_openpyxl_original(label):
    """[S2] §7.3 delivered file is the openpyxl original, not a LibreOffice re-save."""
    p = SAMPLE if label == "Sample" else BLANK
    app = xml_field(app_xml(p), "Application") or ""
    assert not re.search(r"libreoffice|openoffice", app, re.I), f"Application={app!r}"


def test_T00_SS_09_recalc_reproducible(recalculated):
    """[S2] §7.2 reproducibility: two independent recalculations give identical values."""
    a = all_values(recalculated("Sample", SAMPLE, run=1))
    b = all_values(recalculated("Sample", SAMPLE, run=2))
    diff = [k for k in set(a) | set(b) if a.get(k) != b.get(k)]
    assert not diff, f"{len(diff)} cells differ, e.g. {sorted(diff)[:10]}"


# ================================================================== 2.3 EDGE + FIT-05 + OOS-02
SCEN = cases.scenarios()
SCEN_IDS = {
    "FIT-05a_5x6": "T00-FIT-05", "FIT-05b_20x12": "T00-FIT-05", "FIT-05c_allcats": "T00-FIT-05",
    "EDGE-01_blank_No": "T00-EDGE-01", "EDGE-01b_blank_Yes": "T00-EDGE-01",
    "EDGE-02_missing": "T00-EDGE-02", "EDGE-03_excused": "T00-EDGE-03", "EDGE-04_all_tied": "T00-EDGE-04",
    "EDGE-05_one_student": "T00-EDGE-05", "EDGE-06_nothing_due_all_blank": "T00-EDGE-06",
    "OOS-02_capacity_40x60": "T00-OOS-02",
}


@pytest.fixture(scope="session")
def scenario_results(require_lo, lo_tmp):
    cache = {}

    def get(name):
        if name not in cache:
            src = cases.build(BLANK, lo_tmp / "cases" / name / "in.xlsx", **SCEN[name])
            cache[name] = recalc(src, lo_tmp / "cases" / name / "out")
        return cache[name]

    return get


def test_T00_CASE_build_inputs_land_in_place(tmp_path):
    """[harness] scenario inputs are written into the Blank copy where intended (runs without LibreOffice)."""
    out = cases.build(BLANK, tmp_path / "x.xlsx", **SCEN["EDGE-02_missing"])
    wb = openpyxl.load_workbook(out)
    assert wb["Roster"]["B6"].value == cases.names(3)[0]
    assert wb["Assignments"]["B6"].value == "Item 1" and wb["Assignments"]["D8"].value == "Tests"
    assert wb["Gradebook"]["N11"].value == "M"  # student 2, assignment 2
    assert str(wb["Gradebook"]["C10"].value).startswith("=")  # formulas kept
    assert BLANK.stat().st_size > 0 and not (PRODUCT / "x.xlsx").exists()


@pytest.mark.parametrize("name", list(SCEN), ids=[f"{SCEN_IDS[n]}:{n}" for n in SCEN])
def test_T00_EDGE_and_normal_cases(scenario_results, name):
    """[S1] §7.3 edge cases / §7.2 normal+boundary cases: 0 errors and 0 recomputation mismatches."""
    p = scenario_results(name)
    errs = error_cells(p)
    wb = openpyxl.load_workbook(p, data_only=True)
    mism = recompute.compare_workbook(wb)
    extra = []
    res = recompute.class_metrics(recompute.read_model(wb))
    st = wb["Standings"]
    if name == "EDGE-04_all_tied":
        ranks = {st.cell(8 + i, 7).value for i in res}
        pcts = {st.cell(8 + i, 8).value for i in res}
        if ranks != {1}:
            extra.append(f"all tied: ranks {ranks}")
        if len(pcts) != 1:
            extra.append(f"all tied: percentiles differ {pcts}")
    if name == "EDGE-05_one_student" and st.cell(8, 7).value != 1:
        extra.append(f"one student: rank {st.cell(8, 7).value!r}")
    if name == "EDGE-06_nothing_due_all_blank":
        gb = wb["Gradebook"]
        for i in res:
            if recompute.norm(gb.cell(10 + i, 3).value) is not None:
                extra.append(f"nothing graded but overall={gb.cell(10 + i, 3).value!r}")
    if name == "OOS-02_capacity_40x60":
        last = recompute.norm(st.cell(8 + 39, 2).value)
        if last != cases.names(40)[39]:
            extra.append(f"40th student not in Standings: {last!r}")
        if recompute.norm(st.cell(8 + 39, 7).value) is None:
            extra.append("40th student has no rank")
    assert not errs and not mism and not extra, {"errors": errs[:10], "mismatches": mism[:15], "extra": extra}


# ================================================================== 2.4 OOS
def test_T00_OOS_01_limits_documented(workbooks):
    """[S2] §7.2 Start Here states numeric limits for students and assignments."""
    t = start_here_text(workbooks["Sample"])
    stu = re.search(r"(?:up to|more than|max(?:imum)?)\D{0,20}?(\d+)\)?\s*(?:students)?", t, re.I)
    has_students = re.search(r"\b\d+\s+students\b|students\s*\(up to \d+\)", t, re.I)
    has_asg = re.search(r"\b\d+\s+assignments\b|assignments\s*\(up to \d+\)", t, re.I)
    assert stu and has_students and has_asg, "Start Here must state student and assignment limits with numbers"


def test_T00_OOS_03_out_of_range_score_flagged(workbooks):
    """[S2] §7.2 out-of-range scores (> points, negative, stray text) are flagged by validation or CF."""
    wb = workbooks["Blank"]
    gb = wb["Gradebook"]
    ok_dv = [dv for dv in gb.data_validations.dataValidation
             if "M10" in dv.sqref and dv.type in ("decimal", "whole", "custom")]
    ok_cf = []
    for cf in gb.conditional_formatting:
        if "M10" not in cf.sqref:
            continue
        for rule in cf.rules:
            s = " ".join(rule.formula or []).upper()
            if re.search(r"M\$7|<\s*0|ISTEXT|NOT\(ISNUMBER", s):
                ok_cf.append(s)
    assert ok_dv or ok_cf, "no data validation or conditional-format warning for out-of-range scores in Gradebook M10:BT49"


@pytest.mark.skip(reason="MANUAL: T00-OOS-04 41st student / 61st assignment shows as out of range — 未実行")
def test_T00_OOS_04_over_limit_manual():
    """[S2]"""


# ================================================================== 2.5 USE
@BOOKS
def test_T00_USE_01_start_here_first_and_active(workbooks, label):
    """[S3] Start Here is the first tab and the active tab on open."""
    wb = workbooks[label]
    assert wb.sheetnames[0] == "Start Here" and wb.active.title == "Start Here", (wb.sheetnames[0], wb.active.title)


@BOOKS
def test_T00_USE_02_internal_links_resolve(workbooks, label):
    """[S2] §7.2 no broken internal hyperlinks."""
    wb = workbooks[label]
    bad = []
    for ws in wb.worksheets:
        for h in ws._hyperlinks:
            loc = h.location or ""
            if loc:
                target = loc.split("!")[0].strip("'")
                if target not in wb.sheetnames:
                    bad.append((ws.title, h.ref, loc))
    for sheet, coord, f in formulas(wb):
        for m in re.finditer(r'HYPERLINK\(\s*"#\'?([^\'!"]+)', f, re.I):
            if m.group(1) not in wb.sheetnames:
                bad.append((sheet, coord, m.group(0)))
    assert not bad, bad


@BOOKS
def test_T00_USE_03_no_cjk_or_leftovers(workbooks, label):
    """[S3] English product: no CJK text, no TODO/FIXME/lorem/[未検証]/XXX/TBD leftovers."""
    bad = []
    for ws in workbooks[label].worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (CJK.search(c.value) or LEFTOVER.search(c.value)):
                    bad.append((ws.title, c.coordinate, c.value[:60]))
    assert not bad, bad[:20]


@pytest.mark.skip(reason="MANUAL: T00-USE-04 open in Excel and Google Sheets, check truncation — 未実行")
def test_T00_USE_04_open_render_manual():
    """[S2]"""


def test_T00_USE_05_status_words_match_start_here(workbooks):
    """[S3] (added at implementation) Start Here promises assignment statuses 'Submitted, Missing, Excused or
    Not due yet'; the Student Report status column must be able to show exactly those words."""
    wb = workbooks["Sample"]
    t = start_here_text(wb)
    m = re.search(r"marked ([A-Za-z ,]+?)(?:\.|$)", t, re.M)
    assert m, "Start Here does not list the Student Report status words"
    promised = {w.strip() for w in re.split(r",| or ", m.group(1)) if w.strip()}
    lits = set()
    rp = wb["Student Report"]
    for row in rp.iter_rows(min_row=15, max_row=74, min_col=8, max_col=8):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                lits |= {s.strip('"') for s in STR_LIT.findall(c.value)}
    missing = sorted(w for w in promised if w not in lits)
    assert not missing, f"promised {sorted(promised)}; report can show {sorted(x for x in lits if x)}; missing {missing}"


# ================================================================== 2.6 RC
@BOOKS
def test_T00_RC_01_no_hidden_sheets(workbooks, label):
    """[S2] (S0 if secrets inside) no hidden/veryHidden sheets."""
    hidden = [(ws.title, ws.sheet_state) for ws in workbooks[label].worksheets if ws.sheet_state != "visible"]
    assert not hidden, hidden


@BOOKS
def test_T00_RC_02_no_macros_or_external_links(label):
    """[S1] no VBA, external workbook links, data connections or external hyperlinks."""
    p = SAMPLE if label == "Sample" else BLANK
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        bad = [n for n in names if n.endswith("vbaProject.bin") or "externalLink" in n or n.endswith("connections.xml")]
        for n in names:
            if n.endswith(".rels") and 'TargetMode="External"' in z.read(n).decode("utf-8", "replace"):
                bad.append(f"external target in {n}")
    assert not bad, bad


@BOOKS
def test_T00_RC_03_doc_properties_no_personal_data(label):
    """[S0] document properties contain no e-mail or the owner's name."""
    p = SAMPLE if label == "Sample" else BLANK
    with zipfile.ZipFile(p) as z:
        core = z.read("docProps/core.xml").decode("utf-8", "replace")
        app = z.read("docProps/app.xml").decode("utf-8", "replace")
    fields = {t: xml_field(core, t) for t in ("creator", "lastModifiedBy", "title", "subject", "keywords", "description")}
    fields.update({t: xml_field(app, t) for t in ("Company", "Manager")})
    bad = {k: v for k, v in fields.items() if v and (re.search(r"@", v) or re.search(r"\bshun\b", v, re.I))}
    assert not bad, bad


@BOOKS
def test_T00_RC_04_no_contact_details_in_cells(workbooks, label):
    """[S0] no e-mail addresses, phone numbers or street addresses in cell text."""
    pats = [re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+"),
            re.compile(r"\(?\b\d{3}\)?[-. ]\d{3}[-. ]\d{4}\b"),
            re.compile(r"\b\d{1,5}\s+\w+(?:\s\w+)?\s+(?:Street|St\.|Avenue|Ave\.?|Road|Rd\.|Blvd|Lane|Ln\.|Drive|Dr\.)\b")]
    bad = []
    for ws in workbooks[label].worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    if any(p.search(c.value) for p in pats):
                        bad.append((ws.title, c.coordinate, c.value[:60]))
    assert not bad, bad


@pytest.mark.skip(reason="MANUAL: T00-RC-05 sample student names are obviously fictional — 未実行")
def test_T00_RC_05_fictional_names_manual():
    """[S0]"""


# ================================================================== 2.7 SELL (§7.4)
def test_T00_SELL_01_three_linked_sheets(workbooks):
    """[S2 暫定, 較正待ち] §7.4 #1: >= 3 sheets hold formulas that reference other sheets."""
    linked = {s for s, _, f in formulas(workbooks["Sample"]) if refs_other_sheets(f, s)}
    assert len(linked) >= 3, sorted(linked)


@pytest.mark.skip(reason="MANUAL: T00-SELL-01M buyer cannot rebuild in 1 hour [較正待ち] — 未実行")
def test_T00_SELL_01M_manual():
    """[-]"""


@pytest.mark.skip(reason="MANUAL: T00-SELL-02 price band evidence from top listings (no Etsy auto-fetch) [較正待ち] — 未実行")
def test_T00_SELL_02_price_band_manual():
    """[S2 暫定]"""


@pytest.mark.skip(reason="MANUAL: T00-SELL-03 (manual part) sample looks like real data [較正待ち] — 未実行; count part = T00-LST-02")
def test_T00_SELL_03_manual():
    """[S2 暫定]"""


@pytest.mark.skip(reason="MANUAL: T00-SELL-04 (manual part) replace with own data within 5 min [較正待ち] — 未実行")
def test_T00_SELL_04_manual():
    """[S2 暫定]"""


def test_T00_SELL_05a_not_single_sheet(workbooks):
    """[S1] §7.4 #5 (確定): not a single-function one-sheet sheet like v1."""
    with_formulas = {s for s, _, _ in formulas(workbooks["Sample"])}
    assert len(with_formulas) >= 2, sorted(with_formulas)


def test_T00_SELL_05b_price_not_in_v1_band(listing):
    """[S1] §7.4 #5 (確定): price is not in the v1 $1-7 band."""
    need_listing(listing)
    price = listing.get("price_usd")
    assert isinstance(price, (int, float)) and price > 7, price


# ================================================================== 2.8 LST
LISTING_KEYS = {"title": str, "price_usd": (int, float), "tags": list, "description_file": str,
                "faq_file": str, "previews": list, "files_delivered": list, "ai_disclosure": str}


def test_T00_LST_01_listing_json_contract(listing):
    """[S2] listing.json exists with all 8 keys of the right type."""
    need_listing(listing)
    bad = {k: type(listing.get(k)).__name__ for k, t in LISTING_KEYS.items()
           if k not in listing or not isinstance(listing[k], t) or isinstance(listing[k], bool)}
    assert not bad, f"missing/wrong-typed keys: {bad}"


def _png_ok(p: Path) -> str | None:
    b = p.read_bytes()
    if b[:8] != b"\x89PNG\r\n\x1a\n":
        return "bad signature"
    if len(b) < 33 or b[12:16] != b"IHDR":
        return "no IHDR"
    w, h = struct.unpack(">II", b[16:24])
    if w == 0 or h == 0:
        return "zero size"
    if zlib.crc32(b[12:29]) & 0xFFFFFFFF != struct.unpack(">I", b[29:33])[0]:
        return "IHDR CRC mismatch"
    return None


def test_T00_LST_02_five_valid_png_previews(listing):
    """[S2] §7.4 #3: >= 5 preview images, all existing, valid PNG."""
    need_listing(listing)
    prev = listing.get("previews") or []
    problems = []
    for rel in prev:
        p = resolve_listing_path(rel) if isinstance(rel, str) else None
        if p is None:
            problems.append((rel, "missing"))
        elif (e := _png_ok(p)):
            problems.append((rel, e))
    assert len(prev) >= 5 and not problems, {"count": len(prev), "problems": problems}


def test_T00_LST_03_price_in_delegated_band(listing):
    """[S2] §9.3 delegated band $12-19 (outside -> not publishable under delegation)."""
    need_listing(listing)
    price = listing.get("price_usd")
    assert isinstance(price, (int, float)) and 12 <= price <= 19, price


def test_T00_LST_04_ai_disclosure(listing, description_text):
    """[S1] §7.2/§16 description says AI was used to make the product; ai_disclosure text is in it."""
    need_text(listing, description_text)
    disc = (listing.get("ai_disclosure") or "").strip()
    d = " ".join(description_text.split())
    sentences = re.split(r"(?<=[.!?])\s+|\n", description_text)
    explicit = [s for s in sentences if re.search(r"\bAI\b", s) and
                re.search(r"\b(created|made|generated|assisted|used|built|help)", s, re.I)]
    assert disc, "ai_disclosure empty"
    assert explicit, "no explicit AI-use statement in description"
    assert " ".join(disc.split()) in d, "ai_disclosure text not found verbatim in description"


def test_T00_LST_05_description_scope(listing, description_text):
    """[S2] §7.2 description states format, contents (Sample+Blank), limits, updates, returns."""
    need_text(listing, description_text)
    t = description_text
    checks = {
        "format": r"\.xlsx|\bExcel\b|Google Sheets",
        "sample": r"\bsample\b",
        "blank": r"\bblank\b",
        "limits": r"up to \d+|\blimit|not include|does not|\bmax(imum)?\b",
        "updates": r"\bupdate",
        "returns": r"\brefund|\breturn|\bexchange",
    }
    missing = [k for k, pat in checks.items() if not re.search(pat, t, re.I)]
    assert not missing, missing


NUM_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8,
             "nine": 9, "ten": 10, "eleven": 11, "twelve": 12}
TAB_STOP = {"excel", "google", "spreadsheet", "each", "every", "one", "this", "the", "a", "an", "nine",
            "separate", "blank", "sample", "your", "any", "new", "linked", "all"}


def test_T00_LST_06_claimed_tabs_exist(workbooks, listing, description_text):
    """[S1] §7.2 every tab named in the description exists; any 'N tabs' count matches."""
    need_text(listing, description_text)
    wb = workbooks["Sample"]
    real = {norm_tab(ws.title) for ws in visible(wb)}
    bad = []
    for m in re.finditer(r"\b([A-Z][A-Za-z]+(?: [A-Z][A-Za-z]+){0,2}) (?:tab|sheet)s?\b", description_text):
        name = m.group(1)
        words = name.split()
        while words and words[0].lower() in TAB_STOP:
            words = words[1:]
        if not words:
            continue
        cand = norm_tab(" ".join(words))
        if not any(cand == r or cand.endswith(" " + r) or r.startswith(cand) for r in real):
            bad.append(name)
    for m in re.finditer(r"\b(\d+|" + "|".join(NUM_WORDS) + r")[\s-]+(?:[a-z-]+\s+)?(tabs|sheets)\b",
                         description_text, re.I):
        n = int(m.group(1)) if m.group(1).isdigit() else NUM_WORDS[m.group(1).lower()]
        if n != len(real):
            bad.append(f"count claim {m.group(0)!r} vs {len(real)} tabs")
    assert not bad, bad


def _capacity(wb) -> dict:
    ro, asg, pa, su = wb["Roster"], wb["Assignments"], wb["Pacing"], wb["Setup"]
    cnt = lambda ws, col, r0: sum(1 for r in range(r0, ws.max_row + 1) if isinstance(ws.cell(r, col).value, int))
    cats = 0
    r = 12
    while isinstance(su.cell(r, 1).value, str) and su.cell(r, 1).value.strip().lower() != "total":
        cats += 1
        r += 1
    weeks = sum(1 for c in range(8, pa.max_column + 1) if isinstance(pa.cell(4, c).value, int))
    return {"student": cnt(ro, 1, 6), "assignment": cnt(asg, 1, 6), "unit": cnt(pa, 1, 6),
            "categor": cats, "week": weeks}


def test_T00_LST_07_capacity_claims_consistent(workbooks, listing, description_text):
    """[S1] §7.2 numeric capacity claims ('up to N students' ...) are true of the workbook and match Start Here."""
    need_text(listing, description_text)
    wb = workbooks["Blank"]
    cap = _capacity(wb)
    sh = start_here_text(wb)
    bad = []
    for m in re.finditer(r"(?:up to|max(?:imum)?(?: of)?)\s+(\d+)\s+(students?|assignments?|units?|categor(?:y|ies)|weeks?)",
                         description_text, re.I):
        n, noun = int(m.group(1)), m.group(2).lower()
        key = next(k for k in cap if noun.startswith(k))
        if n > cap[key]:
            bad.append(f"{m.group(0)!r} but workbook holds {cap[key]}")
        sh_nums = {int(x) for x in re.findall(rf"(\d+)\s+{key}", sh, re.I)} | \
                  {int(x) for x in re.findall(rf"{key}\w*\s*\(up to (\d+)\)", sh, re.I)}
        if sh_nums and n not in sh_nums:
            bad.append(f"{m.group(0)!r} vs Start Here {sorted(sh_nums)}")
    assert not bad, {"capacity": cap, "problems": bad}


def test_T00_LST_08_files_delivered_hash_and_origin(listing, record_property):
    """[S1] §7.3/§11 delivered files exist, include Sample+Blank, equal the product files, sha256 recorded,
    not a LibreOffice re-save."""
    need_listing(listing)
    files = listing.get("files_delivered") or []
    problems, hashes = [], {}
    for rel in files:
        p = resolve_listing_path(rel) if isinstance(rel, str) else None
        if p is None:
            problems.append((rel, "missing"))
            continue
        h = hashlib.sha256(p.read_bytes()).hexdigest()
        hashes[rel] = h
        app = xml_field(app_xml(p), "Application") or ""
        if re.search(r"libreoffice|openoffice", app, re.I):
            problems.append((rel, f"LibreOffice re-save: Application={app!r}"))
    record_property("sha256", json.dumps(hashes))
    print("files_delivered sha256:", json.dumps(hashes, indent=1))
    base = {Path(r).name for r in files if isinstance(r, str)}
    for req in (SAMPLE, BLANK):
        if req.name not in base:
            problems.append((req.name, "not in files_delivered"))
        else:
            rel = next(r for r in files if Path(r).name == req.name)
            if hashes.get(rel) and hashes[rel] != hashlib.sha256(req.read_bytes()).hexdigest():
                problems.append((rel, "bytes differ from products/ original"))
    recorded = listing.get("sha256") or listing.get("hashes")
    if isinstance(recorded, dict):
        for rel, h in recorded.items():
            if rel in hashes and hashes[rel] != h:
                problems.append((rel, "recorded sha256 mismatch"))
    assert files and not problems, problems


def test_T00_LST_09_etsy_limits(listing):
    """[S2] Etsy limits [未検証]: <= 13 tags, each <= 20 chars; title <= 140 chars."""
    need_listing(listing)
    tags = listing.get("tags") or []
    title = listing.get("title") or ""
    bad = []
    if len(tags) > 13:
        bad.append(f"{len(tags)} tags")
    bad += [f"tag too long: {t!r}" for t in tags if not isinstance(t, str) or len(t) > 20]
    if not title or len(title) > 140:
        bad.append(f"title length {len(title)}")
    assert not bad, bad


def test_T00_LST_10_description_and_faq_exist(listing):
    """[S2] §9.3 description and FAQ exist and are non-empty."""
    need_listing(listing)
    bad = []
    for key in ("description_file", "faq_file"):
        rel = listing.get(key)
        p = resolve_listing_path(rel) if isinstance(rel, str) else None
        if p is None or not p.read_text(encoding="utf-8").strip():
            bad.append(key)
    assert not bad, bad


def test_T00_LST_11_no_cjk_or_leftovers_in_listing(listing):
    """[S3] listing text has no CJK characters or TODO/lorem/[未検証]/TBD leftovers."""
    need_listing(listing)
    texts = {"title": listing.get("title") or ""}
    for key in ("description_file", "faq_file"):
        rel = listing.get(key)
        p = resolve_listing_path(rel) if isinstance(rel, str) else None
        if p is not None:
            texts[key] = p.read_text(encoding="utf-8")
    bad = [k for k, t in texts.items() if CJK.search(t) or LEFTOVER.search(t)]
    assert not bad, bad


@pytest.mark.skip(reason="MANUAL: T00-LST-12 preview images truthfully depict the product — 未実行")
def test_T00_LST_12_previews_match_product_manual():
    """[S2]"""


@pytest.mark.skip(reason="MANUAL: T00-LST-13 Google Sheets claim requires T00-SS-07 to have passed — 未実行")
def test_T00_LST_13_google_claim_manual():
    """[S1]"""
